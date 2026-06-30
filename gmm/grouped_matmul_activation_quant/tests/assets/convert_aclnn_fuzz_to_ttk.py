#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Convert GroupedMatmulActivationQuant aclnn_fuzz cases to TTK csv cases.

Example:
    python gmm/grouped_matmul_activation_quant/tests/assets/convert_aclnn_fuzz_to_ttk.py \
        menkan.xslx --sheet net -o net_ttk.csv
"""

import argparse
import ast
import csv
import math
from pathlib import Path


TTK_COLUMNS = [
    "testcase_name",
    "op_name",
    "stc_inputs",
    "stc_ori_inputs",
    "stc_outputs",
    "stc_ori_outputs",
    "dyn_inputs",
    "dyn_ori_inputs",
    "dyn_outputs",
    "dyn_ori_outputs",
    "stc_input_formats",
    "dyn_input_formats",
    "dyn_input_ori_formats",
    "stc_input_ori_formats",
    "output_formats",
    "output_ori_formats",
    "stc_input_dtypes",
    "dyn_input_dtypes",
    "output_dtypes",
    "other_compilation_params",
    "dyn_input_ranges",
    "dyn_output_ranges",
    "dyn_input_as_list_distribution",
    "stc_input_as_list_distribution",
    "precision_tolerances",
    "input_data_ranges",
    "strict_precision_mode",
    "absolute_precision",
    "shape_check",
    "is_enabled",
]

OP_NAME = "grouped_matmul_activation_quant"
ACLNN_NAME = "aclnnGroupedMatmulActivationQuantWeightNz"
DEFAULT_INPUT_FILE = Path("menkan.xslx")
DEFAULT_OUTPUT_FILE = Path("net_ttk.csv")
DEFAULT_SHEET_NAME = "net"
INPUT_NUM = 6
OUTPUT_NUM = 2
DEFAULT_PRECISION_TOLERANCES = ((0.01, 0.01), (0.01, 0.01))
DEFAULT_INPUT_DATA_RANGES = ((1, 1), (1, 1), (1, 1), (1, 1), (1, 1), (1, 1))
DEFAULT_ABSOLUTE_PRECISION = "1"
DEFAULT_SHAPE_CHECK = "TRUE"
DEFAULT_IS_ENABLED = "1"
DEFAULT_STRICT_PRECISION_MODE = "TRUE"

Y_DTYPE_TO_NAME = {
    35: "float8_e5m2",
    36: "float8_e4m3fn",
}
Y_NAME_TO_DTYPE = {value: key for key, value in Y_DTYPE_TO_NAME.items()}


def parse_literal(value):
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        if text == "":
            return None
        try:
            return ast.literal_eval(text)
        except (ValueError, SyntaxError):
            if text == "True":
                return True
            if text == "False":
                return False
            return text
    return value


def as_tuple(value):
    if value is None:
        return ()
    if isinstance(value, tuple):
        return tuple(as_tuple(item) if isinstance(item, list) else item for item in value)
    if isinstance(value, list):
        return tuple(as_tuple(item) if isinstance(item, list) else item for item in value)
    return value


def ceil_div(left, right):
    return int(math.ceil(int(left) / int(right)))


def dtype_short_name(dtype):
    if dtype == "float8_e4m3fn":
        return "e4"
    if dtype == "float8_e5m2":
        return "e5"
    return str(dtype).replace("float8_", "").replace("float", "f")


def normalize_attr_name(name):
    if name == "turing_config":
        return "tuning_config"
    return name


def normalize_bool(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        return value.strip().lower() == "true"
    return bool(value)


def normalize_int(value, default=0):
    if value is None:
        return default
    return int(value)


def normalize_float(value, default=0.0):
    if value is None:
        return default
    return float(value)


def read_xlsx_rows(path, sheet_name=None):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for reading .xlsx files. Install it with: pip install openpyxl") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    headers = [str(cell).strip() if cell is not None else "" for cell in next(iterator)]

    for values in iterator:
        row = {headers[index]: values[index] if index < len(values) else None for index in range(len(headers))}
        if any(value is not None and str(value).strip() for value in row.values()):
            yield row


def read_csv_rows(path):
    with open(path, newline="", encoding="utf-8-sig") as csv_file:
        yield from csv.DictReader(csv_file)


def read_rows(path, sheet_name=None):
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm", ".xslx"):
        yield from read_xlsx_rows(path, sheet_name)
    elif suffix == ".csv":
        yield from read_csv_rows(path)
    else:
        raise ValueError(f"Unsupported input file type: {path.suffix}. Use .xlsx, .xlsm, or .csv.")


def collect_attrs(row):
    attrs = {}
    for column, raw_name in row.items():
        if not column.startswith("attr_name"):
            continue
        name = parse_literal(raw_name)
        if name is None:
            continue
        suffix = column[len("attr_name"):]
        value_column = "attr_value" + suffix
        if value_column not in row:
            continue
        attrs[normalize_attr_name(str(name))] = parse_literal(row[value_column])
    return attrs


def get_shape_and_dtype(row):
    input_shapes = parse_literal(row.get("input_tensor_shape"))
    input_dtypes = parse_literal(row.get("input_tensor_dtype"))
    output_shapes = parse_literal(row.get("output_tensor_shape"))
    output_dtypes = parse_literal(row.get("output_tensor_dtype"))

    if not isinstance(input_shapes, list) or len(input_shapes) != 6:
        raise ValueError(f"input_tensor_shape must contain 6 entries, got: {input_shapes}")
    if not isinstance(input_dtypes, list) or len(input_dtypes) != 6:
        raise ValueError(f"input_tensor_dtype must contain 6 entries, got: {input_dtypes}")
    if not isinstance(output_shapes, list) or len(output_shapes) != 2:
        raise ValueError(f"output_tensor_shape must contain 2 entries, got: {output_shapes}")
    if not isinstance(output_dtypes, list) or len(output_dtypes) != 2:
        raise ValueError(f"output_tensor_dtype must contain 2 entries, got: {output_dtypes}")

    return input_shapes, input_dtypes, output_shapes, output_dtypes


def extract_mkne(row, input_shapes, attrs):
    x_shape = input_shapes[0]
    weight_shape = input_shapes[2][0]
    transpose_weight = normalize_bool(attrs.get("transpose_weight", False))

    m_value = normalize_int(row.get("M"), x_shape[0])
    k_value = normalize_int(row.get("K"), x_shape[1])
    e_value = normalize_int(row.get("E"), weight_shape[0])
    if row.get("N") not in (None, ""):
        n_value = normalize_int(row.get("N"))
    else:
        n_value = weight_shape[1] if transpose_weight else weight_shape[2]

    return m_value, k_value, n_value, e_value


def weight_nz_shape(e_value, k_value, n_value, transpose_weight):
    if transpose_weight:
        return (e_value, ceil_div(k_value, 32), ceil_div(n_value, 16), 16, 32)
    return (e_value, ceil_div(n_value, 32), ceil_div(k_value, 16), 16, 32)


def normalize_bias_shape(input_shapes):
    bias_shape = input_shapes[4]
    if bias_shape in (None, [], ()):
        return ()
    raise ValueError(f"MXFP8 converted TTK cases expect empty bias, got bias shape: {bias_shape}")


def build_attributes(row_attrs, output_dtypes):
    y_dtype = row_attrs.get("y_dtype")
    if y_dtype is None:
        y_dtype = Y_NAME_TO_DTYPE.get(output_dtypes[0])
    if y_dtype is None:
        raise ValueError(f"Cannot infer y_dtype from output dtype: {output_dtypes[0]}")

    tuning_config = row_attrs.get("tuning_config", [0])
    if tuning_config in (None, []):
        tuning_config = [0]

    return {
        "activation_type": row_attrs.get("activation_type", "gelu_tanh"),
        "transpose_weight": normalize_bool(row_attrs.get("transpose_weight", False)),
        "group_list_type": normalize_int(row_attrs.get("group_list_type", 0)),
        "tuning_config": tuning_config,
        "quant_mode": row_attrs.get("quant_mode", "MX"),
        "y_dtype": normalize_int(y_dtype),
        "round_mode": row_attrs.get("round_mode", "rint"),
        "scale_alg": normalize_int(row_attrs.get("scale_alg", 0)),
        "dst_type_max": normalize_float(row_attrs.get("dst_type_max", 0.0)),
        "group_list_expect": row_attrs.get("group_list", []),
    }


def convert_row(row, row_index):
    aclnn_name = row.get("aclnn_name")
    if aclnn_name and str(aclnn_name).strip() != ACLNN_NAME:
        raise ValueError(f"Unsupported aclnn_name at row {row_index}: {aclnn_name}")

    row_attrs = collect_attrs(row)
    input_shapes, input_dtypes, output_shapes, output_dtypes = get_shape_and_dtype(row)
    attrs = build_attributes(row_attrs, output_dtypes)
    m_value, k_value, n_value, e_value = extract_mkne(row, input_shapes, attrs)

    x_shape = tuple(input_shapes[0])
    group_list_shape = tuple(input_shapes[1])
    weight_logical_shape = tuple(input_shapes[2][0])
    weight_scale_shape = tuple(input_shapes[3][0])
    bias_shape = normalize_bias_shape(input_shapes)
    x_scale_shape = tuple(input_shapes[5])
    y_shape = tuple(output_shapes[0])
    y_scale_shape = tuple(output_shapes[1])
    weight_storage_shape = weight_nz_shape(e_value, k_value, n_value, attrs["transpose_weight"])

    weight_dtype = input_dtypes[2][0]
    weight_scale_dtype = input_dtypes[3][0]
    y_dtype_name = Y_DTYPE_TO_NAME.get(attrs["y_dtype"], output_dtypes[0])
    transpose_tag = "t" if attrs["transpose_weight"] else "nt"
    testcase_name = (
        f"gmaq_aclnn_m{m_value}_k{k_value}_n{n_value}_e{e_value}_"
        f"{transpose_tag}_g{attrs['group_list_type']}_s{attrs['scale_alg']}_"
        f"x{dtype_short_name(input_dtypes[0])}_y{dtype_short_name(y_dtype_name)}"
    )

    stc_inputs = (x_shape, group_list_shape, weight_storage_shape, weight_scale_shape, bias_shape, x_scale_shape)
    stc_ori_inputs = (x_shape, group_list_shape, weight_logical_shape, weight_scale_shape, bias_shape, x_scale_shape)
    stc_outputs = (y_shape, y_scale_shape)
    stc_input_formats = ("ND", "ND", "FRACTAL_NZ", "ND", "ND", "ND")
    stc_input_ori_formats = ("ND", "ND", "ND", "ND", "ND", "ND")
    output_formats = ("ND", "ND")
    stc_input_dtypes = (input_dtypes[0], input_dtypes[1], weight_dtype, weight_scale_dtype, "float32", input_dtypes[5])
    output_dtypes_tuple = (output_dtypes[0], output_dtypes[1])
    dyn_placeholders = tuple((-2,) for _ in range(INPUT_NUM))
    dyn_output_placeholders = tuple((-2,) for _ in range(OUTPUT_NUM))
    as_list_distribution = (0, 0, 1, 1, 1, 0)

    return {
        "testcase_name": testcase_name,
        "op_name": OP_NAME,
        "stc_inputs": repr(stc_inputs),
        "stc_ori_inputs": repr(stc_ori_inputs),
        "stc_outputs": repr(stc_outputs),
        "stc_ori_outputs": repr(stc_outputs),
        "dyn_inputs": repr(dyn_placeholders),
        "dyn_ori_inputs": repr(dyn_placeholders),
        "dyn_outputs": repr(dyn_output_placeholders),
        "dyn_ori_outputs": repr(dyn_output_placeholders),
        "stc_input_formats": repr(stc_input_formats),
        "dyn_input_formats": repr(stc_input_formats),
        "dyn_input_ori_formats": repr(stc_input_ori_formats),
        "stc_input_ori_formats": repr(stc_input_ori_formats),
        "output_formats": repr(output_formats),
        "output_ori_formats": repr(output_formats),
        "stc_input_dtypes": repr(stc_input_dtypes),
        "dyn_input_dtypes": repr(stc_input_dtypes),
        "output_dtypes": repr(output_dtypes_tuple),
        "other_compilation_params": repr(attrs),
        "dyn_input_ranges": "",
        "dyn_output_ranges": "",
        "dyn_input_as_list_distribution": repr(as_list_distribution),
        "stc_input_as_list_distribution": repr(as_list_distribution),
        "precision_tolerances": repr(DEFAULT_PRECISION_TOLERANCES),
        "input_data_ranges": repr(DEFAULT_INPUT_DATA_RANGES),
        "strict_precision_mode": DEFAULT_STRICT_PRECISION_MODE,
        "absolute_precision": DEFAULT_ABSOLUTE_PRECISION,
        "shape_check": DEFAULT_SHAPE_CHECK,
        "is_enabled": DEFAULT_IS_ENABLED,
    }


def convert_file(input_path, output_path, sheet_name=None, limit=None):
    converted_rows = []
    for row_index, row in enumerate(read_rows(input_path, sheet_name), start=2):
        if limit is not None and len(converted_rows) >= limit:
            break
        if not row.get("aclnn_name") and not row.get("input_tensor_shape"):
            continue
        converted_rows.append(convert_row(row, row_index))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w", newline="", encoding="utf-8") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=TTK_COLUMNS)
        writer.writeheader()
        writer.writerows(converted_rows)

    return len(converted_rows)


def main():
    parser = argparse.ArgumentParser(description="Convert GMMAQ aclnn_fuzz xlsx/csv cases to TTK csv format.")
    parser.add_argument(
        "input",
        type=Path,
        nargs="?",
        default=DEFAULT_INPUT_FILE,
        help="Input aclnn_fuzz .xlsx/.xlsm/.csv file. Defaults to menkan.xslx.",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT_FILE,
        help="Output TTK csv path.",
    )
    parser.add_argument("--sheet", default=DEFAULT_SHEET_NAME, help="Worksheet name for .xlsx/.xlsm input.")
    parser.add_argument("--limit", type=int, default=None, help="Convert at most N cases, useful for smoke tests.")
    args = parser.parse_args()

    count = convert_file(args.input, args.output, args.sheet, args.limit)
    print(f"Converted {count} case(s) to {args.output}")


if __name__ == "__main__":
    main()
