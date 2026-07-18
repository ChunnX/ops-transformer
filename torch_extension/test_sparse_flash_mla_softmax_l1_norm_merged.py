# Copyright (c) PyPTO Contributors.
# This program is free software; you can redistribute it and/or modify it under the terms and conditions of
# CANN Open Software License Agreement Version 2.0 (the "License").
# Please refer to the License for details. You may not use this file except in compliance with the License.
# THIS SOFTWARE IS PROVIDED ON AN "AS IS" BASIS, WARRANTIES OF ANY KIND, EITHER EXPRESS OR IMPLIED,
# INCLUDING BUT NOT LIMITED TO NON-INFRINGEMENT, MERCHANTABILITY, OR FITNESS FOR A PARTICULAR PURPOSE.
# See LICENSE in the root of the software repository for full text of the License.
# -----------------------------------------------------------------------------------------------------------

import ast
import logging
import math
import openpyxl
import torch
import torch_npu
import cann_ops_transformer

# ================================================================
#  Constants
# ================================================================
D_TOTAL = 512
SCALE = 1.0 / math.sqrt(D_TOTAL)


def _parse_list(val):
    if val is None or val == "" or val == "no":
        return None
    if isinstance(val, str):
        try:
            return ast.literal_eval(val)
        except:
            return None
    return val


def _to_torch_dtype(dtype_str):
    if dtype_str == "bfloat16":
        return torch.bfloat16
    return torch.float16


def build_cumsum_offsets(a):
    cumsum_a = torch.cumsum(a, dim=0)
    c = torch.cat([torch.zeros(1, dtype=a.dtype, device=a.device), cumsum_a])
    return c.to(a.dtype)


def compress_bnsd_to_tnd(x, seqlist):
    if x is None:
        return None
    B, S, N, D = x.shape
    idxs = torch.arange(S, device=x.device).unsqueeze(0).expand(B, S)
    mask = idxs < seqlist.unsqueeze(1)
    return x[mask]


def compress_bnsd_to_ntd(x, seqlist):
    if x is None:
        return None
    B, N, S, D = x.shape
    x_bsnd = x.transpose(1, 2)
    idxs = torch.arange(S, device=x.device).unsqueeze(0).expand(B, S)
    mask = idxs < seqlist.unsqueeze(1)
    tnd = x_bsnd[mask]
    ntd = tnd.transpose(0, 1)
    return ntd.contiguous()


def generate_seqused(seqlist, b):
    rand_unif = torch.rand(seqlist.shape)
    return (rand_unif * (seqlist.to(torch.float64) + 1)).to(torch.int32)


def generate_sparse_indices_loop(b, sq, skv, k, maskmode, cmp_ratio, layout="BSND", seq_q=None, seq_k=None, has_topk_length=False, cmp_residual_k=None):
    sparse_indices = torch.zeros(b, sq, 1, k, dtype=torch.int32)
    if has_topk_length:
        topk_length = torch.zeros(b, sq, 1, 1, dtype=torch.int32)
    else:
        topk_length = None

    for b_id in range(b):
        cur_sq = sq if layout == "BSND" else seq_q[b_id]
        cur_skv = skv if layout == "BSND" else seq_k[b_id]
        for s1_idx in range(cur_sq):
            if maskmode == 0:
                limit = min(cur_skv, k)
                if has_topk_length:
                    if limit > 0:
                        topk_length[b_id, s1_idx, 0, :] = torch.randint(0, limit, (1,), dtype=torch.int32)
                    else:
                        topk_length[b_id, s1_idx, 0, :] = 0
                    limit = int(topk_length[b_id, s1_idx, 0, 0])
                rand_vals = torch.randint(0, cur_skv, (limit,), dtype=torch.int32)
                sparse_indices[b_id, s1_idx, 0, :limit] = rand_vals
            elif maskmode == 3:
                if cmp_ratio > 1:
                    cur_skv_ = cur_skv * cmp_ratio + int(cmp_residual_k[b_id])
                else:
                    cur_skv_ = cur_skv * cmp_ratio
                actual_sk = max((cur_skv_ - cur_sq + s1_idx + 1) // cmp_ratio, 0)
                limit = min(actual_sk, k)
                if has_topk_length:
                    if limit > 0:
                        topk_length[b_id, s1_idx, 0, :] = torch.randint(0, limit, (1,), dtype=torch.int32)
                    else:
                        topk_length[b_id, s1_idx, 0, :] = 0
                    limit = int(topk_length[b_id, s1_idx, 0, 0])
                if limit > 0:
                    rand_vals = torch.randint(0, actual_sk, (limit,), dtype=torch.int32)
                    sparse_indices[b_id, s1_idx, 0, :limit] = rand_vals
    return sparse_indices, topk_length


def compute_softmax_l1_norm_loop(q, k, softmax_lse, seqused_q=None, seqused_k=None, cmp_residual_k=None, sparse_indices=None, topk_length=None, mask_mode=0, cmp_ratio=1, core_num=32, layout="BSND", seq_q=None, seq_k=None, softmax_scale=SCALE):
    B, Sq, G, _ = q.shape
    Sk = k.shape[1]
    device = q.device

    metadata = torch.zeros(1024, dtype=torch.int32, device=device)
    if seqused_q is not None:
        usedQ = seqused_q
    elif layout == "BSND":
        usedQ = [Sq for _ in range(B)]
    else:
        usedQ = seq_q.numpy()
    total_t1 = sum(usedQ)
    if total_t1 <= core_num:
        used_core_num = total_t1
        remain_core_num = 0
        former_core_process_num = 1
        remain_core_process_num = 0
    else:
        former_core_process_num = math.ceil(total_t1 / core_num)
        remain_core_process_num = former_core_process_num - 1
        former_core_num = total_t1 - core_num * remain_core_process_num
        remain_core_num = core_num - former_core_num
        used_core_num = core_num
    metadata[0] = total_t1
    metadata[1] = former_core_process_num
    metadata[2] = remain_core_process_num
    metadata[3] = remain_core_num
    metadata[4] = used_core_num

    if sparse_indices is not None:
        k_length = sparse_indices.shape[-1]
    else:
        k_length = Sk

    softmax_l1_norm = torch.zeros(B, Sq, 1, k_length, dtype=torch.float32, device=device)
    for b_id in range(B):
        sq_len = int(seqused_q[b_id].item() if seqused_q is not None else (Sq if layout == "BSND" else seq_q[b_id].item()))
        sk_len = int(seqused_k[b_id].item() if seqused_k is not None else (Sk if layout == "BSND" else seq_k[b_id].item()))

        for s1_id in range(sq_len):
            query = q[b_id, s1_id, :, :].float()
            if mask_mode == 0:
                actual_sk = sk_len
            else:
                if cmp_ratio > 1:
                    ori_sk_len = sk_len * cmp_ratio + cmp_residual_k[b_id]
                else:
                    ori_sk_len = sk_len * cmp_ratio
                actual_sk = max((ori_sk_len - sq_len + s1_id + 1) // cmp_ratio, 0)

            if sparse_indices is not None:
                K = sparse_indices.shape[-1]
                if topk_length is not None:
                    K = min(int(topk_length[b_id, s1_id, 0, 0]), K)
                actual_sk = min(actual_sk, K)
                if actual_sk > 0:
                    indices = sparse_indices[b_id, s1_id, 0, :actual_sk]
                    key = k[b_id, indices, 0, :].float()
                else:
                    key = k[b_id, :0, 0, :].float()
            else:
                key = k[b_id, :actual_sk, 0, :].float()

            qk = query @ key.T
            qk_scale = qk * softmax_scale
            lse = softmax_lse[b_id, 0, s1_id, :]
            p = torch.exp(qk_scale - lse.unsqueeze(1))
            softmax_l1 = p.sum(dim=0) / G

            if actual_sk > 0:
                softmax_l1_norm[b_id, s1_id, 0, :actual_sk] = softmax_l1[:actual_sk]

    return softmax_l1_norm, metadata


def find_mismatched_elements(actual, expected, rtol=5e-3, atol=5e-3):
    actual_flat = actual.cpu().flatten().to(torch.float32)
    expected_flat = expected.cpu().flatten().to(torch.float32)
    abs_diff = torch.abs(actual_flat - expected_flat)
    tolerance = atol + rtol * torch.abs(expected_flat)
    mismatch_mask = abs_diff > tolerance
    nan_mask = torch.isnan(actual_flat) | torch.isnan(expected_flat)
    inf_mask = torch.isinf(actual_flat) | torch.isinf(expected_flat)
    mismatch_mask = mismatch_mask | nan_mask | inf_mask
    mismatch_indices = torch.nonzero(mismatch_mask, as_tuple=False).flatten()
    total = actual_flat.numel()
    mismatch_count = mismatch_indices.numel()
    print(f"总元素数: {total}")
    print(f"异常点数: {mismatch_count} ({mismatch_count / total * 100:.4f}%)")
    print("-" * 80)
    if mismatch_count > 0:
        print(f"{'Index':<15} {'Actual':<20} {'Expected':<20} {'AbsDiff':<20} {'Tolerance':<20}")
        print("-" * 80)
        print_count = 50
        for idx in mismatch_indices:
            i = idx.item()
            a = actual_flat[i].item()
            e = expected_flat[i].item()
            d = abs_diff[i].item()
            t = tolerance[i].item()
            print(f"{i:<15} {a:<20.8f} {e:<20.8f} {d:<20.8f} {t:<20.8f}")
            print_count -= 1
            if print_count <= 0:
                break
        print("-" * 80)
    return mismatch_count == 0


def run_test_case(params):
    device = "npu:0"
    torch.npu.set_device(device)
    torch.manual_seed(params.get("manual_seed", 42))

    B = params["B"]
    S1 = params["S1"]
    S2 = params["S2"]
    G = params["N1"]
    D = params["D"]
    dtype = _to_torch_dtype(params.get("dtype", "float16"))
    layout = params.get("layout_q", "BSND")
    mask_mode = params.get("mask_mode", 0)
    cmp_ratio = params.get("cmp_ratio", 1)
    softmax_scale = params.get("softmax_scale", SCALE)
    core_num = 32

    # Sparse info
    K = params.get("sparse_size")
    is_sparse = K is not None and K > 0

    # Topk info
    topk_mode = params.get("topk_mode")
    has_topk_length = topk_mode is not None and str(topk_mode).lower() not in ["no", ""]

    # Extract lists
    seq_q_list = _parse_list(params.get("seqlens_list_q"))
    seq_k_list = _parse_list(params.get("seqlens_list_kv"))
    seqused_q_list = _parse_list(params.get("seqused_q"))
    seqused_k_list = _parse_list(params.get("seqused_kv"))
    cmp_residual_k_list = _parse_list(params.get("cmp_residual_kv"))

    # Resolve actual lengths per batch
    if layout == "TND":
        # TND: S1 and S2 are sums of all batches
        if seq_q_list is None:
            if seqused_q_list is not None:
                seq_q_list = seqused_q_list
            else:
                seq_q_list = [S1 // B] * B
                seq_q_list[-1] += S1 - sum(seq_q_list)
        if seq_k_list is None:
            if seqused_k_list is not None:
                seq_k_list = seqused_k_list
            else:
                seq_k_list = [S2 // B] * B
                seq_k_list[-1] += S2 - sum(seq_k_list)
        
        # Ensure length matches B
        while len(seq_q_list) < B: seq_q_list.append(1)
        seq_q_list = seq_q_list[:B]
        while len(seq_k_list) < B: seq_k_list.append(1)
        seq_k_list = seq_k_list[:B]

        seq_q_cpu = torch.tensor(seq_q_list, dtype=torch.int32)
        seq_k_cpu = torch.tensor(seq_k_list, dtype=torch.int32)
        
        if seqused_q_list is not None:
            seqused_q_cpu = torch.tensor(seqused_q_list, dtype=torch.int32)
            while len(seqused_q_cpu) < B:
                seqused_q_cpu = torch.cat([seqused_q_cpu, torch.randint(1, max(seq_q_list)+1, (B - len(seqused_q_cpu),), dtype=torch.int32)])
        else:
            seqused_q_cpu = None
            
        if seqused_k_list is not None:
            seqused_k_cpu = torch.tensor(seqused_k_list, dtype=torch.int32)
            while len(seqused_k_cpu) < B:
                seqused_k_cpu = torch.cat([seqused_k_cpu, torch.randint(1, max(seq_k_list)+1, (B - len(seqused_k_cpu),), dtype=torch.int32)])
        else:
            seqused_k_cpu = None
            
        gen_s1 = max(seq_q_list)
        gen_s2 = max(seq_k_list)
    else:
        # BSND: S1 and S2 are max lengths
        gen_s1 = S1
        gen_s2 = S2
        
        if seqused_q_list is not None:
            seqused_q_cpu = torch.tensor(seqused_q_list, dtype=torch.int32)
            while len(seqused_q_cpu) < B:
                seqused_q_cpu = torch.cat([seqused_q_cpu, torch.randint(1, S1+1, (B - len(seqused_q_cpu),), dtype=torch.int32)])
        else:
            seqused_q_cpu = None
            
        if seqused_k_list is not None:
            seqused_k_cpu = torch.tensor(seqused_k_list, dtype=torch.int32)
            while len(seqused_k_cpu) < B:
                seqused_k_cpu = torch.cat([seqused_k_cpu, torch.randint(1, S2+1, (B - len(seqused_k_cpu),), dtype=torch.int32)])
        else:
            seqused_k_cpu = None

        seq_q_cpu = torch.full((B,), S1, dtype=torch.int32)
        seq_k_cpu = torch.full((B,), S2, dtype=torch.int32)

    logging.info(f"\nRunning test: {params.get('testcase_name', 'N/A')} (layout={layout}, sparse={is_sparse}, topk={has_topk_length})")

    # Generate base BSND data on CPU first for golden computation
    query = torch.rand((B, gen_s1, G, D), dtype=dtype)
    key = torch.rand((B, gen_s2, 1, D), dtype=dtype)
    lse = torch.rand((B, 1, gen_s1, G), dtype=torch.float32)

    # Cmp residual
    cmp_residual_k_cpu = None
    if cmp_ratio > 1:
        if cmp_residual_k_list is not None:
            cmp_residual_k_cpu = torch.tensor(cmp_residual_k_list, dtype=torch.int32)
            while len(cmp_residual_k_cpu) < B:
                cmp_residual_k_cpu = torch.cat([cmp_residual_k_cpu, torch.randint(0, cmp_ratio, (B - len(cmp_residual_k_cpu),), dtype=torch.int32)])
        else:
            cmp_residual_k_cpu = torch.randint(0, cmp_ratio, (B,), dtype=torch.int32)

    # Sparse indices
    if is_sparse:
        sparse_indices_cpu, topk_length_cpu = generate_sparse_indices_loop(
            B, gen_s1, gen_s2, K, mask_mode, cmp_ratio, "BSND", None, None, has_topk_length, cmp_residual_k_cpu
        )
    else:
        sparse_indices_cpu, topk_length_cpu = None, None

    # Golden
    golden_out, metadata = compute_softmax_l1_norm_loop(
        query, key, lse, seqused_q_cpu, seqused_k_cpu, cmp_residual_k_cpu,
        sparse_indices_cpu, topk_length_cpu, mask_mode, cmp_ratio, core_num,
        "BSND", None, None, softmax_scale
    )

    # Move to device
    query_dev = query.to(device)
    key_dev = key.to(device)
    lse_dev = lse.to(device)
    seqused_q_dev = seqused_q_cpu.to(device) if seqused_q_cpu is not None else None
    seqused_k_dev = seqused_k_cpu.to(device) if seqused_k_cpu is not None else None
    cmp_residual_k_dev = cmp_residual_k_cpu.to(device) if cmp_residual_k_cpu is not None else None
    metadata_dev = metadata.to(device)

    if layout == "TND":
        seq_q_dev = build_cumsum_offsets(seq_q_cpu.to(device))
        seq_k_dev = build_cumsum_offsets(seq_k_cpu.to(device))

        query_tnd = compress_bnsd_to_tnd(query_dev, seq_q_cpu.to(device))
        key_tnd = compress_bnsd_to_tnd(key_dev, seq_k_cpu.to(device))
        lse_tnd = compress_bnsd_to_ntd(lse_dev, seq_q_cpu.to(device))

        if sparse_indices_cpu is not None:
            sparse_indices_tnd = compress_bnsd_to_tnd(sparse_indices_cpu.to(device), seq_q_cpu.to(device))
            if topk_length_cpu is not None:
                topk_length_tnd = compress_bnsd_to_tnd(topk_length_cpu.to(device), seq_q_cpu.to(device)).squeeze(-1)
            else:
                topk_length_tnd = None
        else:
            sparse_indices_tnd = None
            topk_length_tnd = None

        golden_tnd = compress_bnsd_to_tnd(golden_out, seq_q_cpu)

        max_seqlen_k = S2
        npu_out = cann_ops_transformer.sparse_flash_mla_softmax_l1_norm(
            query_tnd, key_tnd, lse_tnd,
            sparse_indices=sparse_indices_tnd,
            cu_seqlens_q=seq_q_dev,
            cu_seqlens_k=seq_k_dev,
            seqused_q=seqused_q_dev,
            seqused_k=seqused_k_dev,
            cmp_residual_k=cmp_residual_k_dev,
            topk_length=topk_length_tnd,
            metadata=metadata_dev,
            scale_value=softmax_scale,
            max_seqlen_k=max_seqlen_k,
            cmp_ratio=cmp_ratio,
            mask_mode=mask_mode,
            layout_q="TND",
            layout_k="TND",
        )

        golden_out_final = golden_tnd
    else:
        # BSND
        sparse_indices_dev = sparse_indices_cpu.to(device) if sparse_indices_cpu is not None else None
        topk_length_dev = topk_length_cpu.to(device) if topk_length_cpu is not None else None
        max_seqlen_k = gen_s2 if not is_sparse else 0

        npu_out = cann_ops_transformer.sparse_flash_mla_softmax_l1_norm(
            query_dev, key_dev, lse_dev,
            sparse_indices=sparse_indices_dev,
            cu_seqlens_q=None,
            cu_seqlens_k=None,
            seqused_q=seqused_q_dev,
            seqused_k=seqused_k_dev,
            cmp_residual_k=cmp_residual_k_dev,
            topk_length=topk_length_dev,
            metadata=metadata_dev,
            scale_value=softmax_scale,
            max_seqlen_k=max_seqlen_k,
            cmp_ratio=cmp_ratio,
            mask_mode=mask_mode,
            layout_q="BSND",
            layout_k="BSND",
        )
        golden_out_final = golden_out

    torch.npu.synchronize()

    print(f"npu_out shape: {npu_out.shape}")
    print(f"golden_out shape: {golden_out_final.shape}")
    diff = (npu_out.cpu().flatten().to(torch.float32) - golden_out_final.cpu().flatten().to(torch.float32)).abs().max().item()
    print(f"Max diff: {diff}")

    is_match = find_mismatched_elements(npu_out, golden_out_final, rtol=5e-3, atol=5e-3)
    if is_match:
        print("✅ 精度校验通过，所有点均在容差范围内。")
    else:
        print("❌ 精度校验失败，存在异常点。")

    torch.testing.assert_close(
        npu_out.cpu().flatten().to(torch.float32),
        golden_out_final.cpu().flatten().to(torch.float32),
        rtol=5e-3, atol=5e-3
    )
    logging.info(f"  PASS: {params.get('testcase_name', 'N/A')}")


def load_excel_params(filepath):
    wb = openpyxl.load_workbook(filepath)
    ws = wb["Sheet1"]
    headers = [cell.value for cell in ws[1]]
    test_cases = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        params = {}
        for i, h in enumerate(headers):
            if h is None:
                continue
            # Normalize header names
            h_clean = h.strip().lower().replace(" ", "_")
            val = row[i]
            if h_clean in ["seqlens_list_q", "seqlens_list_kv", "seqused_q", "seqused_kv", "cmp_residual_kv"]:
                params[h_clean] = _parse_list(val)
            else:
                params[h_clean] = val
        test_cases.append(params)
    return test_cases


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
    logging.info("SparseFlashMlaSoftmaxL1Norm Merged Tests")

    excel_path = r"C:\daiyuwen\01.code\ops-transformer\torch_extension\smla_softmax_l1_norm.xlsx"
    test_cases = load_excel_params(excel_path)

    for idx, case in enumerate(test_cases):
        logging.info("%s", '=' * 60)
        logging.info(f"Test Case {idx+1}: {case.get('testcase_name', 'N/A')}")
        try:
            run_test_case(case)
        except Exception as e:
            logging.error(f"Test Case {idx+1} failed with exception: {e}")
            import traceback
            traceback.print_exc()

    logging.info("\nAll tests completed!")