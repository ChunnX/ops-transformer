# -----------------------------------------------------------------------------------------------------------
# Copyright (c) 2026 Huawei Technologies Co., Ltd.
# This program is free software, you can redistribute it and/or modify it under the terms and conditions of
# CANN Open Software License Agreement Version 2.0 (the "License").
# Please refer to the License for details. You may not use this file except in compliance with the License.
# THIS SOFTWARE IS PROVIDED ON AN "AS IS" BASIS, WARRANTIES OF ANY KIND, EITHER EXPRESS OR IMPLIED,
# INCLUDING BUT NOT LIMITED TO NON-INFRINGEMENT, MERCHANTABILITY, OR FITNESS FOR A PARTICULAR PURPOSE.
# See LICENSE in the root of the software repository for the full text of the License.
# -----------------------------------------------------------------------------------------------------------

"""
Test script for FAG (Flash Attn Grad) metadata verification from xlsx.

Flow:
  1. Parse fag_test_case.xlsx to load test cases.
  2. For each enabled case: call flash_attn_metadata (is_grad_enabled=True).
  3. Extract FAG region, compare with Python reference.
  4. Report pass/fail per case.

Usage:
    python test_flash_attn_grad_metadata.py --xlsx fag_test_case.xlsx
    python test_flash_attn_grad_metadata.py
"""

import ast
import os
import re
import argparse

import torch
import torch_npu
import cann_ops_transformer

try:
    import openpyxl
except ImportError:
    openpyxl = None

# ============================================================
# FAG metadata field indices (must match flash_attn_metadata.h)
# ============================================================
HEAD_FAG_START_OFFSET_INDEX = 4
FAG_SPLIT_AXIS_INDEX = 0
FAG_BLOCK_OUTER_INDEX = 1
FAG_BLOCK_FACTOR_INDEX = 2
FAG_S1_OUTER_INDEX = 3
FAG_S2_OUTER_INDEX = 4
FAG_LAYOUT_TYPE_INDEX = 5
FAG_IS_SPARSE_INDEX = 6
FAG_SPARSE_MODE_INDEX = 7
FAG_BLOCK_STARTS_OFFSET = 8
FAG_BLOCK_ENDS_OFFSET = 44
FAG_TND_START_BIDX_OFFSET = 80
FAG_CORE_LIST_NUM = 36

S1CV_RATIO = 2
S2CV_RATIO = 1

LAYOUT_BSND = 1
LAYOUT_BNSD = 3
LAYOUT_TND = 4

SPARSE_NO_MASK = 0
SPARSE_RIGHT_DOWN_CAUSAL = 3
SPARSE_BAND = 4


def _align_to(value, alignment):
    if alignment <= 0:
        return value
    return (value + alignment - 1) // alignment * alignment


def _ref_dense_split(b, n2, g, s1_outer, s2_outer, aic_num):
    fused_outer = b * n2 * g * s1_outer * s2_outer
    block_factor = (fused_outer + aic_num - 1) // aic_num
    block_outer = (fused_outer + block_factor - 1) // block_factor
    starts = [0] * FAG_CORE_LIST_NUM
    ends = [0] * FAG_CORE_LIST_NUM
    for i in range(block_outer):
        starts[i] = block_factor * i
        ends[i] = min(block_factor * (i + 1), fused_outer)
    return starts, ends, block_outer, block_factor


def _ref_get_parse(s1, s2, s1_cv_inner, cv_s2_inner, s2_outer, s1_token, s2_token):
    parse_info = []
    for i in range(s2_outer):
        lip = max(0, cv_s2_inner * i - s2_token)
        if lip > s1:
            begin = (s1 + s1_cv_inner - 1) // s1_cv_inner
        else:
            begin = lip // s1_cv_inner
        if i == s2_outer - 1:
            cv_block_tail = cv_s2_inner if (s2 % cv_s2_inner == 0) else (s2 % cv_s2_inner)
        else:
            cv_block_tail = cv_s2_inner
        rp = min(max(0, cv_s2_inner * i + cv_block_tail + s1_token), s1)
        end = (rp + s1_cv_inner - 1) // s1_cv_inner
        tmp_size = max(0, end - begin) if end > begin else 0
        parse_info.append((begin, end, tmp_size))
    cum = 0
    for i in range(len(parse_info)):
        cum += parse_info[i][2]
        parse_info[i] = (parse_info[i][0], parse_info[i][1], cum)
    return parse_info


def _ref_check_invalid(parse_info, s1_outer):
    is_invalid_col = False
    is_invalid_row = False
    covered = [False] * s1_outer
    for begin, end, _ in parse_info:
        if begin >= end:
            is_invalid_col = True
        for j in range(begin, min(end, s1_outer)):
            covered[j] = True
    if any(not c for c in covered):
        is_invalid_row = True
    return is_invalid_col, is_invalid_row


def _ref_sparse_block_info(b, n2, g, s1, s2, s1_inner, s2_inner, s1_cv_inner, cv_s2_inner,
                           s1_outer, s2_outer, s1_token, s2_token, aic_num):
    parse_info = _ref_get_parse(s1, s2, s1_cv_inner, cv_s2_inner, s2_outer, s1_token, s2_token)
    s1s2o_count = parse_info[s2_outer - 1][2] if s2_outer > 0 else 0
    fused_outer = b * n2 * g * s1s2o_count
    if fused_outer == 0:
        return [0]*FAG_CORE_LIST_NUM, [0]*FAG_CORE_LIST_NUM, 0, 0
    block_factor = (fused_outer + aic_num - 1) // aic_num
    block_outer = (fused_outer + block_factor - 1) // block_factor
    n2gs = n2 * g * s1s2o_count
    gs = g * s1s2o_count
    starts = [0] * FAG_CORE_LIST_NUM
    ends = [0] * FAG_CORE_LIST_NUM
    starts[0] = 0
    ends[block_outer - 1] = b * n2 * g * s1_outer * s2_outer
    for c in range(1, block_outer):
        current_idx = min(c * block_factor, fused_outer)
        b_idx = current_idx // n2gs
        b_tail = current_idx % n2gs
        n2_idx = b_tail // gs
        n2_tail = b_tail % gs
        g_idx = n2_tail // s1s2o_count
        g_tail = n2_tail % s1s2o_count
        s1o_idx = 0
        s2o_idx = 0
        pre_size = 0
        for i in range(s2_outer):
            next_size = parse_info[i][2]
            if g_tail >= pre_size and g_tail < next_size:
                s2o_idx = i
                s1o_idx = parse_info[i][0] + g_tail - pre_size - 1
                break
            pre_size = next_size
        starts[c] = (((b_idx * n2 + n2_idx) * g + g_idx) * s2_outer + s2o_idx) * s1_outer + s1o_idx + 1
        ends[c - 1] = starts[c]
    return starts, ends, block_outer, block_factor


def _ref_bn2s2_dense_split(b, n2, g, s1_outer, s2_outer, aic_num):
    fused_outer = b * n2 * g * s2_outer
    bns2_factor = (fused_outer + aic_num - 1) // aic_num
    block_outer = (fused_outer + bns2_factor - 1) // bns2_factor
    total_block = fused_outer * s1_outer
    block_factor = bns2_factor * s1_outer
    starts = [0] * FAG_CORE_LIST_NUM
    ends = [0] * FAG_CORE_LIST_NUM
    for i in range(block_outer):
        starts[i] = block_factor * i
        ends[i] = min(block_factor * (i + 1), total_block)
    return starts, ends, block_outer, block_factor


def _ref_bn2_dense_split(b, n2, g, s1_outer, s2_outer, aic_num):
    fused_outer = b * n2 * g
    bn_factor = (fused_outer + aic_num - 1) // aic_num
    block_outer = (fused_outer + bn_factor - 1) // bn_factor
    block_factor = bn_factor * (s1_outer * s2_outer)
    fused_outer *= (s1_outer * s2_outer)
    starts = [0] * FAG_CORE_LIST_NUM
    ends = [0] * FAG_CORE_LIST_NUM
    for i in range(block_outer):
        starts[i] = block_factor * i
        ends[i] = min(block_factor * (i + 1), fused_outer)
    return starts, ends, block_outer, block_factor


def _ref_bn2_sparse_block_info(b, n2, g, s1, s2, s1_inner, s2_inner, s1_cv_inner, cv_s2_inner,
                                s1_outer, s2_outer, s1_token, s2_token, aic_num):
    parse_info = _ref_get_parse(s1, s2, s1_cv_inner, cv_s2_inner, s2_outer, s1_token, s2_token)
    s1s2o_count = parse_info[s2_outer - 1][2] if s2_outer > 0 else 0
    fused_outer = b * n2 * g
    bn_factor = (fused_outer + aic_num - 1) // aic_num
    block_outer = (fused_outer + bn_factor - 1) // bn_factor
    fused_outer *= s1s2o_count
    block_factor = bn_factor * s1s2o_count
    n2gs = n2 * g * s1s2o_count
    gs = g * s1s2o_count
    starts = [0] * FAG_CORE_LIST_NUM
    ends = [0] * FAG_CORE_LIST_NUM
    starts[0] = 0
    ends[block_outer - 1] = b * n2 * g * s1_outer * s2_outer
    for c in range(1, block_outer):
        current_idx = min(c * block_factor, fused_outer)
        b_idx = current_idx // n2gs
        b_tail = current_idx % n2gs
        n2_idx = b_tail // gs
        n2_tail = b_tail % gs
        g_idx = n2_tail // s1s2o_count
        g_tail = n2_tail % s1s2o_count
        s1o_idx = 0
        s2o_idx = 0
        pre_size = 0
        for i in range(s2_outer):
            next_size = parse_info[i][2]
            if g_tail >= pre_size and g_tail < next_size:
                s2o_idx = i
                s1o_idx = parse_info[i][0] + g_tail - pre_size - 1
                break
            pre_size = next_size
        starts[c] = (((b_idx * n2 + n2_idx) * g + g_idx) * s2_outer + s2o_idx) * s1_outer + s1o_idx + 1
        ends[c - 1] = starts[c]
    return starts, ends, block_outer, block_factor


def _ref_tnd_load_balance(b, n2, g, seqlens_q, seqlens_kv,
                          s1_cv_inner, cv_s2_inner, s1_outer, s2_outer,
                          is_sparse, s1_token, s2_token, aic_num):
    batch_info = []
    for i in range(b):
        actual_s1 = seqlens_q[i]
        actual_s2 = seqlens_kv[i]
        actual_s1o = (actual_s1 + s1_cv_inner - 1) // s1_cv_inner
        actual_s2o = (actual_s2 + cv_s2_inner - 1) // cv_s2_inner
        if not is_sparse:
            valid_blocks = actual_s1o * actual_s2o
            parse = None
        else:
            parse = _ref_get_parse(actual_s1, actual_s2, s1_cv_inner, cv_s2_inner,
                                   actual_s2o, s1_token, s2_token)
            valid_blocks = parse[actual_s2o - 1][2] if actual_s2o > 0 else 0
        batch_info.append((actual_s1o, actual_s2o, valid_blocks, parse))

    cum_valid = [0] * b
    cum_full = [0] * b
    for i in range(b):
        n2g_valid = n2 * g * batch_info[i][2]
        n2g_full = n2 * g * batch_info[i][0] * batch_info[i][1]
        cum_valid[i] = (cum_valid[i - 1] + n2g_valid) if i > 0 else n2g_valid
        cum_full[i] = (cum_full[i - 1] + n2g_full) if i > 0 else n2g_full

    fused_outer = cum_valid[b - 1] if b > 0 else 0
    if fused_outer == 0:
        return [0]*FAG_CORE_LIST_NUM, [0]*FAG_CORE_LIST_NUM, 0, 0, [0]*FAG_CORE_LIST_NUM

    block_factor = (fused_outer + aic_num - 1) // aic_num
    block_outer = (fused_outer + block_factor - 1) // block_factor

    starts = [0] * FAG_CORE_LIST_NUM
    ends = [0] * FAG_CORE_LIST_NUM
    tnd_start_bidx = [0] * FAG_CORE_LIST_NUM
    starts[0] = 0
    ends[block_outer - 1] = cum_full[b - 1]

    for c in range(1, block_outer):
        current_idx = min(c * block_factor, fused_outer)
        b_idx = 0
        for bb in range(b):
            if cum_valid[bb] > current_idx:
                b_idx = bb
                break
        tnd_start_bidx[c] = b_idx
        b_tail = current_idx if b_idx == 0 else current_idx - cum_valid[b_idx - 1]
        actual_s1o, actual_s2o, valid_blocks, parse = batch_info[b_idx]
        n2gs = g * valid_blocks
        if n2gs == 0:
            n2gs = 1
        n2_idx = b_tail // n2gs
        n2_tail = b_tail % n2gs
        g_idx = n2_tail // valid_blocks if valid_blocks > 0 else 0
        g_tail = n2_tail % valid_blocks if valid_blocks > 0 else 0

        if parse is None:
            s2o_idx = g_tail // actual_s1o if actual_s1o > 0 else 0
            s1o_idx = g_tail % actual_s1o if actual_s1o > 0 else 0
        else:
            s2o_idx = 0
            s1o_idx = 0
            pre_size = 0
            for j in range(actual_s2o):
                next_size = parse[j][2]
                if g_tail >= pre_size and g_tail < next_size:
                    s2o_idx = j
                    s1o_idx = parse[j][0] + g_tail - pre_size
                    break
                pre_size = next_size

        full_blocks = actual_s1o * actual_s2o
        offset = (n2_idx * g + g_idx) * full_blocks + s2o_idx * actual_s1o + s1o_idx
        starts[c] = offset if b_idx == 0 else cum_full[b_idx - 1] + offset
        ends[c - 1] = starts[c]

    return starts, ends, block_outer, block_factor, tnd_start_bidx


def _ref_bn2s2_tnd_split(b, n2, g, seqlens_q, seqlens_kv,
                          s1_cv_inner, cv_s2_inner, s1_outer, s2_outer,
                          sparse_mode, s1_token, s2_token, aic_num):
    S1CV_RATIO_DEFAULT = 2
    FAG_S1CV_RATIO_DEFAULT = 2
    FAG_S2CV_RATIO_DEFAULT = 1
    s1_inner = 64
    s2_inner = 128

    def calcle_actual_token(sm, st1, st2, as1, as2):
        a1, a2 = st1, st2
        if sm == 3 or sm == 4:
            a1 = a1 + as1 - as2
            a2 = a2 - as1 + as2
        return a1, a2

    total_block_info = [[0, 0] for _ in range(b)]
    actural_block_info = [[0.0] * s2_outer for _ in range(b)]
    total_sum = 0.0
    max_unit = 0.0

    for i in range(b):
        as1 = seqlens_q[i]
        as2 = seqlens_kv[i]
        as1o = (as1 + s1_cv_inner - 1) // s1_cv_inner
        as2o = (as2 + cv_s2_inner - 1) // cv_s2_inner
        total_block_info[i][0] = as1o * as2o

        ac_s1t, ac_s2t = calcle_actual_token(sparse_mode, s1_token, s2_token, as1, as2)

        for j in range(s2_outer):
            if cv_s2_inner * j >= as2:
                actural_block_info[i][j] = 0.0
            else:
                lip = max(cv_s2_inner * j - ac_s2t, 0)
                cv_block_tail = as2 - cv_s2_inner * j if cv_s2_inner * (j + 1) > as2 else cv_s2_inner
                actural_s1_begin = float(min(lip, as1) if lip > as1 else lip)
                actural_s1_end = float(min(as1, max(cv_s2_inner * j + cv_block_tail + ac_s1t, 0)))
                actural_s1_num = 0 if actural_s1_begin > actural_s1_end else actural_s1_end - actural_s1_begin
                actural_s2_num = float(cv_block_tail)
                actural_block_info[i][j] = actural_s1_num / float(s1_cv_inner) + actural_s2_num / float(cv_s2_inner)
                total_sum += actural_block_info[i][j] * n2 * g
                max_unit = max(max_unit, actural_block_info[i][j])

        if i == 0:
            total_block_info[0][1] = n2 * g * total_block_info[0][0]
        else:
            total_block_info[i][1] = n2 * g * total_block_info[i][0] + total_block_info[i - 1][1]

    def is_possible(possible_max):
        current = 0.0
        need = 1
        n2g = n2 * g
        for i in range(b * n2g):
            bi = i // n2g
            for j in range(s2_outer):
                num = actural_block_info[bi][j]
                if current + num > possible_max:
                    need += 1
                    current = num
                else:
                    current += num
                if need > aic_num:
                    return False
        return True

    left = max_unit
    right = total_sum
    while left < right - 1:
        mid = (left + right) / 2
        if is_possible(mid):
            right = mid
        else:
            left = mid + 1

    max_per_core = right
    current = 0.0
    core_idx = 0
    starts = [0] * FAG_CORE_LIST_NUM
    ends = [0] * FAG_CORE_LIST_NUM
    tnd_start_bidx = [0] * FAG_CORE_LIST_NUM

    for bb in range(b):
        for n in range(n2 * g):
            as1o = (seqlens_q[bb] + s1_cv_inner - 1) // s1_cv_inner
            for j in range(s2_outer):
                num = actural_block_info[bb][j]
                if core_idx >= FAG_CORE_LIST_NUM:
                    return starts, ends, 0, 0, tnd_start_bidx
                if current + num > max_per_core:
                    pre_batch = total_block_info[bb - 1][1] if bb > 0 else 0
                    pre_ng = n * total_block_info[bb][0]
                    pre_s2 = j * as1o
                    ends[core_idx] = pre_batch + pre_ng + pre_s2
                    starts[core_idx + 1] = ends[core_idx]
                    core_idx += 1
                    current = num
                    tnd_start_bidx[core_idx] = bb
                else:
                    current += num

    starts[0] = 0
    ends[core_idx] = total_block_info[b - 1][1]
    block_outer = core_idx + 1
    block_factor = int(max_per_core + 0.999999)

    for c in range(block_outer, FAG_CORE_LIST_NUM):
        starts[c] = 0
        ends[c] = 0
        tnd_start_bidx[c] = 0

    return starts, ends, block_outer, block_factor, tnd_start_bidx


def _compute_fag_reference(batch_size, num_heads_q, num_heads_kv, head_dim,
                           max_seqlen_q, max_seqlen_kv,
                           mask_mode, win_left, win_right,
                           layout_q, aic_core_num=36,
                           seqlens_q=None, seqlens_kv=None):
    s1 = max_seqlen_q
    s2 = max_seqlen_kv
    n1 = num_heads_q
    n2 = num_heads_kv
    g = n1 // n2
    b = batch_size
    s1_inner = 64
    s2_inner = 128
    s1_cv_inner = min(s1_inner * S1CV_RATIO, s1)
    cv_s2_inner = min(s2_inner * S2CV_RATIO, s2)
    s1_outer = (s1 + s1_cv_inner - 1) // s1_cv_inner
    s2_outer = (s2 + cv_s2_inner - 1) // cv_s2_inner

    if layout_q == "TND":
        layout_type = LAYOUT_TND
    elif layout_q == "BNSD":
        layout_type = LAYOUT_BNSD
    else:
        layout_type = LAYOUT_BSND

    sparse_mode = SPARSE_NO_MASK
    is_sparse = False
    s1_token = win_left if win_left is not None and win_left != -1 else 2147483647
    s2_token = win_right if win_right is not None and win_right != -1 else 2147483647

    if mask_mode == 0:
        is_sparse = False
    elif mask_mode == 3:
        sparse_mode = SPARSE_RIGHT_DOWN_CAUSAL
        is_sparse = True
        s1_token = 2147483647
        s2_token = 0
        if layout_type != LAYOUT_TND:
            s1_token = s1_token + s1 - s2
            s2_token = s2_token - s1 + s2
    elif mask_mode == 4:
        sparse_mode = SPARSE_BAND
        if s1 > s1_token or s2 > s2_token:
            is_sparse = True
        if layout_type != LAYOUT_TND:
            s1_token = s1_token + s1 - s2
            s2_token = s2_token - s1 + s2

    is_tnd = (layout_type == LAYOUT_TND)
    is_all_same = False
    # SupportTransBSND: TND + all batches same seq lengths → degrade to BSND
    if is_tnd:
        if seqlens_q is not None and seqlens_kv is not None:
            is_all_same = all(s == seqlens_q[0] for s in seqlens_q) and \
                          all(s == seqlens_kv[0] for s in seqlens_kv)
        else:
            is_all_same = True
    use_tnd_split = is_tnd
    if is_all_same and is_tnd:
        layout_type = LAYOUT_BSND
        use_tnd_split = False

    is_bn2 = (s1 <= 128 and s2 <= 128) and (n1 == n2) and (head_dim <= 512)
    bn_limit = ((b * n1) >= 256) or ((b * n1) >= 128 and (s1 % 128 == 0) and (s2 % 128 == 0))
    bn_sparse_limit = bn_limit and not use_tnd_split
    is_bn2_multiblk = bn_sparse_limit and (s1 > 128 or s2 > 128) and (s1 <= 640 and s2 <= 640) and (n1 == n2) and (head_dim <= 512)
    is_bn2 = is_bn2_multiblk if is_bn2_multiblk else is_bn2
    if is_bn2 and not is_bn2_multiblk:
        if use_tnd_split and head_dim > 128:
            is_bn2 = False

    bn2s2_not_tnd_limit = (s1 < s2) and (s2 <= 1024) and (s2 - s1 >= 128) and (head_dim <= 128) and not is_sparse
    bn2s2_route_limit = (head_dim <= 512) and (use_tnd_split or is_all_same or bn2s2_not_tnd_limit) and (n1 == n2)

    if is_bn2:
        split_axis = 1
    elif bn2s2_route_limit:
        split_axis = 2
        if is_all_same:
            use_tnd_split = True
    else:
        split_axis = 0

    starts = [0] * FAG_CORE_LIST_NUM
    ends = [0] * FAG_CORE_LIST_NUM
    block_outer = 0
    block_factor = 0

    # Layer 1: BN2S2
    if split_axis == 2:
        if is_sparse or use_tnd_split:
            if is_all_same:
                starts, ends, block_outer, block_factor = _ref_bn2s2_dense_split(
                    b, n2, g, s1_outer, s2_outer, aic_core_num)
                tnd_start_bidx_s2 = [0] * FAG_CORE_LIST_NUM
                for c in range(1, block_outer):
                    current_idx = min(c * block_factor, b * n2 * g * s2_outer)
                    tnd_start_bidx_s2[c] = min(current_idx // (n2 * g * s2_outer), b - 1)
            else:
                starts, ends, block_outer, block_factor, tnd_start_bidx_s2 = _ref_bn2s2_tnd_split(
                    b, n2, g, seqlens_q, seqlens_kv,
                    s1_cv_inner, cv_s2_inner, s1_outer, s2_outer,
                    sparse_mode, s1_token, s2_token, aic_core_num)
        else:
            starts, ends, block_outer, block_factor = _ref_bn2s2_dense_split(b, n2, g, s1_outer, s2_outer, aic_core_num)
            tnd_start_bidx_s2 = [0] * FAG_CORE_LIST_NUM
        if block_outer >= aic_core_num:
            result = {
                'split_axis': split_axis, 'block_outer': block_outer, 'block_factor': block_factor,
                's1_outer': s1_outer, 's2_outer': s2_outer, 'layout_type': layout_type,
                'is_sparse': 1 if is_sparse else 0, 'sparse_mode': sparse_mode,
                'block_starts': starts, 'block_ends': ends,
                'tnd_start_bidx': tnd_start_bidx_s2,
            }
            return result
        # BN2S2 didn't produce enough cores, fall through

    # Layer 2: BN2 MultiBlk
    if split_axis == 1 and is_bn2_multiblk:
        if is_sparse:
            starts, ends, block_outer, block_factor = _ref_bn2_sparse_block_info(
                b, n2, g, s1, s2, s1_inner, s2_inner, s1_cv_inner, cv_s2_inner,
                s1_outer, s2_outer, s1_token, s2_token, aic_core_num)
        else:
            starts, ends, block_outer, block_factor = _ref_bn2_dense_split(b, n2, g, s1_outer, s2_outer, aic_core_num)
        s1s2o_ok = True
        if is_sparse:
            parse_info = _ref_get_parse(s1, s2, s1_cv_inner, cv_s2_inner, s2_outer, s1_token, s2_token)
            s1s2o_count = parse_info[s2_outer - 1][2] if s2_outer > 0 else 0
            is_invalid_col, is_invalid_row = _ref_check_invalid(parse_info, s1_outer)
            s1s2o_ok = (s1s2o_count > 0) and (not is_invalid_col) and (not is_invalid_row)
        if s1s2o_ok:
            result = {
                'split_axis': split_axis, 'block_outer': block_outer, 'block_factor': block_factor,
                's1_outer': s1_outer, 's2_outer': s2_outer, 'layout_type': layout_type,
                'is_sparse': 1 if is_sparse else 0, 'sparse_mode': sparse_mode,
                'block_starts': starts, 'block_ends': ends,
                'tnd_start_bidx': [0] * FAG_CORE_LIST_NUM,
            }
            return result
        # Invalid: degrade
        is_bn2 = False
        is_bn2_multiblk = False
        split_axis = 0

    # Layer 3: fallback — splitAxis = isBn2 ? BN2 : BN2GS1S2
    split_axis = 1 if is_bn2 else 0
    tnd_start_bidx = [0] * FAG_CORE_LIST_NUM
    if use_tnd_split:
        starts, ends, block_outer, block_factor, tnd_start_bidx = _ref_tnd_load_balance(
            b, n2, g, seqlens_q, seqlens_kv,
            s1_cv_inner, cv_s2_inner, s1_outer, s2_outer,
            is_sparse, s1_token, s2_token, aic_core_num)
    elif is_sparse:
        starts, ends, block_outer, block_factor = _ref_sparse_block_info(
            b, n2, g, s1, s2, s1_inner, s2_inner, s1_cv_inner, cv_s2_inner,
            s1_outer, s2_outer, s1_token, s2_token, aic_core_num)
    elif is_bn2:
        starts, ends, block_outer, block_factor = _ref_bn2_dense_split(b, n2, g, s1_outer, s2_outer, aic_core_num)
    else:
        starts, ends, block_outer, block_factor = _ref_dense_split(b, n2, g, s1_outer, s2_outer, aic_core_num)

    return {
        'split_axis': split_axis,
        'block_outer': block_outer,
        'block_factor': block_factor,
        's1_outer': s1_outer,
        's2_outer': s2_outer,
        'layout_type': layout_type,
        'is_sparse': 1 if is_sparse else 0,
        'sparse_mode': sparse_mode,
        'block_starts': starts,
        'block_ends': ends,
        'tnd_start_bidx': tnd_start_bidx,
    }


def _extract_fag_from_metadata(metadata_tensor):
    meta = metadata_tensor.cpu().tolist()
    fag_offset = meta[HEAD_FAG_START_OFFSET_INDEX]
    if fag_offset == 0:
        return None
    fag = meta[fag_offset:]
    return {
        'split_axis': fag[FAG_SPLIT_AXIS_INDEX],
        'block_outer': fag[FAG_BLOCK_OUTER_INDEX],
        'block_factor': fag[FAG_BLOCK_FACTOR_INDEX],
        's1_outer': fag[FAG_S1_OUTER_INDEX],
        's2_outer': fag[FAG_S2_OUTER_INDEX],
        'layout_type': fag[FAG_LAYOUT_TYPE_INDEX],
        'is_sparse': fag[FAG_IS_SPARSE_INDEX],
        'sparse_mode': fag[FAG_SPARSE_MODE_INDEX],
        'block_starts': fag[FAG_BLOCK_STARTS_OFFSET:FAG_BLOCK_STARTS_OFFSET + FAG_CORE_LIST_NUM],
        'block_ends': fag[FAG_BLOCK_ENDS_OFFSET:FAG_BLOCK_ENDS_OFFSET + FAG_CORE_LIST_NUM],
        'tnd_start_bidx': fag[FAG_TND_START_BIDX_OFFSET:FAG_TND_START_BIDX_OFFSET + FAG_CORE_LIST_NUM],
        'raw_fag_offset': fag_offset,
    }


def _print_fag_verbose(label, actual, expected):
    print(f"\n  --- {label} verbose ---")
    print(f"  splitAxis:    actual={actual['split_axis']}  expected={expected['split_axis']}")
    print(f"  blockOuter:   actual={actual['block_outer']}  expected={expected['block_outer']}")
    print(f"  blockFactor:  actual={actual['block_factor']}  expected={expected['block_factor']}")
    print(f"  s1Outer:      actual={actual['s1_outer']}  expected={expected['s1_outer']}")
    print(f"  s2Outer:      actual={actual['s2_outer']}  expected={expected['s2_outer']}")
    print(f"  layoutType:   actual={actual['layout_type']}  expected={expected['layout_type']}")
    print(f"  isSparse:     actual={actual['is_sparse']}  expected={expected['is_sparse']}")
    print(f"  sparseMode:   actual={actual['sparse_mode']}  expected={expected['sparse_mode']}")
    bo = actual['block_outer']
    print(f"  blockStarts (first {bo} active cores):")
    for i in range(min(bo, FAG_CORE_LIST_NUM)):
        a_val = actual['block_starts'][i]
        e_val = expected['block_starts'][i]
        match = "OK" if a_val == e_val else "MISMATCH"
        print(f"    blockStarts[{i:2d}] = {a_val:12d}  (expected={e_val:12d})  {match}")
    print(f"  blockEnds (first {bo} active cores):")
    for i in range(min(bo, FAG_CORE_LIST_NUM)):
        a_val = actual['block_ends'][i]
        e_val = expected['block_ends'][i]
        match = "OK" if a_val == e_val else "MISMATCH"
        print(f"    blockEnds[{i:2d}]   = {a_val:12d}  (expected={e_val:12d})  {match}")
    if bo < FAG_CORE_LIST_NUM:
        all_zero_start = all(actual['block_starts'][i] == 0 for i in range(bo, FAG_CORE_LIST_NUM))
        all_zero_end = all(actual['block_ends'][i] == 0 for i in range(bo, FAG_CORE_LIST_NUM))
        print(f"  unused cores [{bo}..{FAG_CORE_LIST_NUM-1}]: starts_all_zero={all_zero_start} ends_all_zero={all_zero_end}")
    print(f"  --- end verbose ---\n")


# ============================================================
# Old FAG KJCTEST summary parsing (for three-way comparison)
# ============================================================

_KJCTEST_BLOCK_PAIR_RE = re.compile(
    r'blockStarts\[(\d+)\]\s*=\s*(\d+).*?blockEnds\[(\d+)\]\s*=\s*(\d+)'
)
_SPLIT_AXIS_RE = re.compile(r'splitAxis\[(\d+)\]')


def _parse_old_fag_summary(summary_path):
    """Parse kjctest_summary.txt into {case_name: old_fag_result_dict}."""
    cases = {}
    current_case = None
    current_lines = []

    with open(summary_path, 'r') as f:
        for line in f:
            line = line.rstrip()
            if line.startswith('=== Case:'):
                if current_case is not None:
                    cases[current_case] = _parse_kjctest_lines(current_lines)
                current_case = line.replace('=== Case:', '').strip().rstrip('===').strip()
                current_lines = []
            elif current_case is not None and not line.startswith('#'):
                current_lines.append(line)

    if current_case is not None:
        cases[current_case] = _parse_kjctest_lines(current_lines)

    return cases


def _parse_kjctest_lines(lines):
    """Extract split info from a list of KJCTEST + GetTilingKey lines."""
    result = {
        'split_axis': None,
        'block_starts': [0] * FAG_CORE_LIST_NUM,
        'block_ends': [0] * FAG_CORE_LIST_NUM,
        'block_outer': 0,
        'route': None,
        'raw_lines': lines,
    }

    for line in lines:
        if 'GetTilingKey' in line and 'splitAxis' in line:
            m = _SPLIT_AXIS_RE.search(line)
            if m:
                result['split_axis'] = int(m.group(1))

        if 'KJCTEST' in line:
            if 'bn2multiblk' in line:
                result['route'] = 'bn2multiblk'
            elif 'bn2gs1s2 or bn2' in line:
                result['route'] = 'bn2gs1s2_or_bn2'
            elif 'bn2gs1s2' in line:
                result['route'] = 'bn2gs1s2'

            m_pair = _KJCTEST_BLOCK_PAIR_RE.search(line)
            if m_pair:
                idx = int(m_pair.group(1))
                start = int(m_pair.group(2))
                end = int(m_pair.group(4))
                if idx < FAG_CORE_LIST_NUM:
                    result['block_starts'][idx] = start
                    result['block_ends'][idx] = end

    block_outer = 0
    for i in range(FAG_CORE_LIST_NUM):
        if result['block_starts'][i] != 0 or result['block_ends'][i] != 0:
            block_outer = i + 1
    result['block_outer'] = block_outer

    return result


def _print_three_way_verbose(name, actual, expected, old_fag):
    """Print three-way comparison: new metadata vs Python reference vs old FAG."""
    print(f"\n  --- Three-way comparison: {name} ---")

    labels = ['splitAxis', 'blockOuter', 'blockFactor', 's1Outer', 's2Outer',
              'layoutType', 'isSparse', 'sparseMode']
    fields = ['split_axis', 'block_outer', 'block_factor', 's1_outer', 's2_outer',
              'layout_type', 'is_sparse', 'sparse_mode']

    print(f"  {'Field':<14s} {'New(meta)':>12s} {'Py(ref)':>12s} {'Old(FAG)':>12s}  Status")
    for label, field in zip(labels, fields):
        a = actual[field]
        e = expected[field]
        o = old_fag.get(field)
        o_str = str(o) if o is not None else 'N/A'
        if o is not None and a == e == o:
            status = 'ALL_MATCH'
        elif o is not None and a != o:
            status = 'NEW!=OLD'
        elif o is not None and e != o:
            status = 'REF!=OLD'
        elif a != e:
            status = 'NEW!=REF'
        else:
            status = 'OK'
        print(f"  {label:<14s} {str(a):>12s} {str(e):>12s} {o_str:>12s}  {status}")

    bo = actual['block_outer']
    print(f"  blockStarts (first {bo} active cores):")
    for i in range(min(bo, FAG_CORE_LIST_NUM)):
        a_val = actual['block_starts'][i]
        e_val = expected['block_starts'][i]
        o_val = old_fag['block_starts'][i] if old_fag['block_outer'] > i else None
        o_str = str(o_val) if o_val is not None else 'N/A'
        parts = []
        if a_val == e_val:
            parts.append('new==ref')
        else:
            parts.append(f'NEW!=REF(a={a_val},e={e_val})')
        if o_val is not None:
            if a_val == o_val:
                parts.append('new==old')
            else:
                parts.append(f'NEW!=OLD(o={o_val})')
        print(f"    [{i:2d}] new={a_val:12d} ref={e_val:12d} old={o_str:>12s}  {' '.join(parts)}")

    print(f"  blockEnds (first {bo} active cores):")
    for i in range(min(bo, FAG_CORE_LIST_NUM)):
        a_val = actual['block_ends'][i]
        e_val = expected['block_ends'][i]
        o_val = old_fag['block_ends'][i] if old_fag['block_outer'] > i else None
        o_str = str(o_val) if o_val is not None else 'N/A'
        parts = []
        if a_val == e_val:
            parts.append('new==ref')
        else:
            parts.append(f'NEW!=REF(a={a_val},e={e_val})')
        if o_val is not None:
            if a_val == o_val:
                parts.append('new==old')
            else:
                parts.append(f'NEW!=OLD(o={o_val})')
        print(f"    [{i:2d}] new={a_val:12d} ref={e_val:12d} old={o_str:>12s}  {' '.join(parts)}")

    print(f"  old FAG route: {old_fag.get('route', 'N/A')}")
    print(f"  --- end three-way ---\n")




_DTYPE_FILTER = {"float16", "bfloat16", "float32"}


def _parse_cell(val):
    if val is None:
        return None
    if isinstance(val, bool):
        return val
    if isinstance(val, str):
        s = val.strip()
        if s.lower() == "null" or s == "":
            return None
        if s.startswith("[") and s.endswith("]"):
            return ast.literal_eval(s)
        try:
            return int(s)
        except ValueError:
            try:
                return float(s)
            except ValueError:
                return s
    return val


def _load_xlsx_cases(xlsx_path):
    if openpyxl is None:
        raise ImportError("openpyxl is required. Install via: pip install openpyxl")
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    cases = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        case = {}
        for i, h in enumerate(headers):
            if h is None or i >= len(row):
                continue
            case[h] = _parse_cell(row[i])
        enable_val = case.get("enable")
        if enable_val is False or (isinstance(enable_val, str) and enable_val.lower() == "false"):
            continue
        dtype_str = case.get("dtype", "float16")
        if dtype_str not in _DTYPE_FILTER:
            continue
        cases.append(case)
    return cases


def _run_and_check(case, old_fag_data=None):
    name = case.get("testcase_name", "unnamed")
    B = case["B"]
    N1 = case["N1"]
    N2 = case["N2"]
    D = case["D"]
    S1 = case["S1"]
    S2 = case["S2"]
    layout = case.get("input_layout", "BSND") or "BSND"
    mask_mode = case.get("sparse_mode", 0) or 0
    win_left = case.get("win_left")
    win_right = case.get("win_right")
    wl = win_left if win_left is not None else -1
    wr = win_right if win_right is not None else -1

    print(f"\n  case: {name} (B={B} N1={N1} N2={N2} S1={S1} S2={S2} D={D}"
          f" mask={mask_mode} layout={layout})")

    cu_seqlens_q = None
    cu_seqlens_kv = None
    seqlens_q = case.get("seqlens_list_q")
    seqlens_kv = case.get("seqlens_list_kv")
    if layout == "TND":
        if seqlens_q is None:
            seqlens_q = [S1] * B
        if seqlens_kv is None:
            seqlens_kv = [S2] * B
        cu_q = [0]
        for s in seqlens_q:
            cu_q.append(cu_q[-1] + s)
        cu_kv = [0]
        for s in seqlens_kv:
            cu_kv.append(cu_kv[-1] + s)
        cu_seqlens_q = torch.tensor(cu_q, dtype=torch.int32, device="npu")
        cu_seqlens_kv = torch.tensor(cu_kv, dtype=torch.int32, device="npu")

    metadata = cann_ops_transformer.flash_attn_metadata(
        num_heads_q=N1,
        num_heads_kv=N2,
        head_dim=D,
        cu_seqlens_q=cu_seqlens_q,
        cu_seqlens_kv=cu_seqlens_kv,
        batch_size=B,
        max_seqlen_q=S1,
        max_seqlen_kv=S2,
        mask_mode=mask_mode,
        win_left=wl,
        win_right=wr,
        layout_q=layout,
        layout_kv=layout,
        layout_out=layout,
        is_grad_enabled=True,
    )
    assert metadata is not None, "metadata tensor is None"
    assert metadata.dtype == torch.int32, f"metadata dtype={metadata.dtype}, expected int32"

    actual = _extract_fag_from_metadata(metadata)
    assert actual is not None, "FAG region not populated (head[4]==0)"

    expected = _compute_fag_reference(
        B, N1, N2, D, S1, S2, mask_mode, wl, wr, layout,
        seqlens_q=seqlens_q, seqlens_kv=seqlens_kv,
    )

    int_fields = ['split_axis', 'block_outer', 'block_factor', 's1_outer', 's2_outer',
                  'layout_type', 'is_sparse', 'sparse_mode']
    failures = []
    for field in int_fields:
        a = actual[field]
        e = expected[field]
        if a != e:
            failures.append((field, f"got={a} expected={e}"))

    for field in ['block_starts', 'block_ends', 'tnd_start_bidx']:
        a_list = actual[field]
        e_list = expected[field]
        for i in range(FAG_CORE_LIST_NUM):
            if a_list[i] != e_list[i]:
                failures.append((f"{field}[{i}]", f"got={a_list[i]} expected={e_list[i]}"))

    if old_fag_data is not None:
        old_fag = old_fag_data.get(name)
        if old_fag is None:
            print(f"  [WARN] No old FAG data for case: {name}")
            _print_fag_verbose("MATCHED (no old data)", actual, expected)
        elif old_fag['split_axis'] is None:
            print(f"  [WARN] Old FAG data for case '{name}' has no splitAxis (incomplete KJCTEST?)")
            _print_fag_verbose("MATCHED (no old data)", actual, expected)
        else:
            _print_three_way_verbose(name, actual, expected, old_fag)

    if failures:
        msg = "\n".join(f"  FAIL {f}: {d}" for f, d in failures)
        _print_fag_verbose("ACTUAL (mismatch)", actual, expected)
        raise AssertionError(f"FAG metadata mismatch:\n{msg}")

    if old_fag_data is None:
        _print_fag_verbose("MATCHED", actual, expected)


def main():
    parser = argparse.ArgumentParser(description="FAG metadata verification from xlsx")
    parser.add_argument("--xlsx", type=str, default=None,
                        help="Path to xlsx test case file")
    parser.add_argument("--compare-old-fag", action="store_true",
                        help="Enable three-way comparison: new metadata vs Python reference vs old FAG KJCTEST logs")
    parser.add_argument("--old-fag-summary", type=str, default=None,
                        help="Path to kjctest_summary.txt (generated by test_fusion_attn_grad_v2.py --collect-kjctest)")
    args = parser.parse_args()

    if args.xlsx is None:
        script_dir = os.path.dirname(os.path.abspath(__file__))
        args.xlsx = os.path.join(script_dir, "fag_test_case.xlsx")

    print(f"[fuzz] xlsx: {args.xlsx}")

    old_fag_data = None
    if args.compare_old_fag:
        summary_path = args.old_fag_summary
        if summary_path is None:
            summary_path = os.path.join(os.path.dirname(os.path.abspath(args.xlsx)), "kjctest_summary.txt")
        if not os.path.exists(summary_path):
            print(f"[fuzz] ERROR: old FAG summary not found: {summary_path}")
            print(f"[fuzz] Generate it first: python test_fusion_attn_grad_v2.py --collect-kjctest --log-file <log>")
            return
        old_fag_data = _parse_old_fag_summary(summary_path)
        print(f"[fuzz] Three-way comparison enabled, loaded {len(old_fag_data)} old FAG cases from: {summary_path}")

    cases = _load_xlsx_cases(args.xlsx)
    print(f"[fuzz] Loaded {len(cases)} cases (enabled + fp16/bf16/fp32)")

    pass_count = 0
    fail_count = 0
    for idx, case in enumerate(cases):
        name = case.get("testcase_name", f"case_{idx}")
        try:
            _run_and_check(case, old_fag_data)
            torch.npu.synchronize()
            pass_count += 1
        except Exception as e:
            print(f"  [FAIL] {name}: {type(e).__name__}: {e}")
            fail_count += 1
            torch.npu.synchronize()

    print(f"\n{'='*60}")
    print(f"  Summary: {pass_count} passed, {fail_count} failed / {len(cases)} total")
    if old_fag_data is not None:
        print(f"  Three-way comparison details printed above per case")
    print(f"{'='*60}")


if __name__ == "__main__":
    main()
