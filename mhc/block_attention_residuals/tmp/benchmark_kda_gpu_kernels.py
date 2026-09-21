#!/usr/bin/env python3
"""Benchmark selected KDA/GLA Triton kernels on CUDA GPU.

The benchmark cases are loaded from tools/KDA算子启用10个Case汇总.xlsx by
default. The workbook contains 10 cases for each listed operator; every case is
run only against the operator named in its row.

Results are written to an Excel workbook (default: tools/benchmark_results.xlsx).

Examples:
    python tools/benchmark_kda_gpu_kernels.py --device 0
    python tools/benchmark_kda_gpu_kernels.py --device 0 --iters 100 --warmup 20
    python tools/benchmark_kda_gpu_kernels.py --dry-run
    python tools/benchmark_kda_gpu_kernels.py --op-filter chunk_kda_bwd_dAv --case-filter 01,shape7
    python tools/benchmark_kda_gpu_kernels.py --device 0 --output-xlsx /tmp/kda_results.xlsx
    python tools/benchmark_kda_gpu_kernels.py --device 0 --no-save
    python tools/benchmark_kda_gpu_kernels.py --device 0 --op-filter chunk_kda_bwd_intra --no-safe-gate
"""

from __future__ import annotations

import argparse
import os
import re
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Callable
from xml.etree import ElementTree as ET

os.environ.setdefault("TORCH_DEVICE_BACKEND_AUTOLOAD", "0")

torch = None
triton = None


def load_gpu_modules() -> None:
    global torch, triton
    if torch is not None and triton is not None:
        return
    import torch as torch_module
    import triton as triton_module

    torch = torch_module
    triton = triton_module


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CASE_XLSX = REPO_ROOT / "tools" / "KDA算子启用10个Case汇总.xlsx"
DEFAULT_OUTPUT_XLSX = REPO_ROOT / "tools" / "benchmark_results.xlsx"
sys.path.insert(0, str(REPO_ROOT))

OPS = (
    "chunk_gla_fwd_kernel_o",
    "chunk_kda_bwd_dAv",
    "chunk_kda_bwd_intra",
    "chunk_kda_bwd_wy_dqkg_fused",
    "chunk_kda_fwd_intra_inter_solve_fused",
    "chunk_kda_fwd_intra_intra_sub_chunk",
    "chunk_kda_fwd_intra_token_parrallel",
)
TOKEN_PARALLEL_OP = "chunk_kda_fwd_intra_token_parrallel"
INTER_SOLVE_OP = "chunk_kda_fwd_intra_inter_solve_fused"


@dataclass(frozen=True)
class BenchCase:
    op: str
    op_index: int
    label: str
    layout: str
    B: int
    T: int
    H: int
    HV: int
    K: int
    V: int
    BT: int = 64
    BC: int = 16
    dtype: str = "float16"
    safe_gate: bool | None = None
    disable_recompute: bool = False
    cu_seqlens: tuple[int, ...] | None = None
    config: str = ""
    source_row: int = 0

    @property
    def case_id(self) -> str:
        return f"{self.op_index:02d}:{self.label}"

    @property
    def is_varlen(self) -> bool:
        return self.layout.lower().startswith("var") or self.cu_seqlens is not None

    def shape_desc(self) -> str:
        desc = (
            f"{self.layout},B={self.B},T={self.T},H={self.H},HV={self.HV},"
            f"K={self.K},V={self.V},BT={self.BT},BC={self.BC},{self.dtype}"
        )
        if self.safe_gate is not None:
            desc += f",safe_gate={self.safe_gate}"
        if self.cu_seqlens is not None:
            desc += f",seqs={len(self.cu_seqlens) - 1}"
        return desc


@dataclass
class BenchResult:
    op: str
    case_id: str
    shape: str
    mean_ms: float
    p50_ms: float
    p90_ms: float
    min_ms: float
    max_ms: float
    note: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--device", type=int, default=0)
    parser.add_argument("--warmup", type=int, default=10)
    parser.add_argument("--iters", type=int, default=50)
    parser.add_argument("--seed", type=int, default=0)
    parser.add_argument("--case-xlsx", type=Path, default=DEFAULT_CASE_XLSX)
    parser.add_argument("--output-xlsx", type=Path, default=DEFAULT_OUTPUT_XLSX,
                        help=f"Path to save benchmark results as an Excel workbook (default: {DEFAULT_OUTPUT_XLSX}).")
    parser.add_argument("--no-save", action="store_true",
                        help="Do not save results to an Excel/CSV file.")
    parser.add_argument("--dry-run", action="store_true", help="Only load and print cases; do not require CUDA or import FLA kernels.")
    parser.add_argument("--safe-gate", action="store_true", help="Override xlsx safe_gate with True for supported KDA operators.")
    parser.add_argument("--no-safe-gate", action="store_true",
                        help="Force safe_gate=False for all cases. Workaround for the Triton 3.2.0 "
                             "'b_gn bf16 vs fp32' CompilationError in chunk_kda_bwd_kernel_intra.")
    parser.add_argument("--single-case", action="store_true", help="Run only the shape specified by --op/--B/--T/--H/--HV/--K/--V.")
    parser.add_argument("--op", choices=OPS, default="chunk_kda_bwd_dAv", help="Operator used with --single-case.")
    parser.add_argument("--case-filter", default="", help="Comma-separated case ids/labels/op indexes to run, for example 01,shape7.")
    parser.add_argument("--op-filter", default="", help="Comma-separated op names to run. Empty means all ops.")
    parser.add_argument("--B", type=int, default=2)
    parser.add_argument("--T", type=int, default=128)
    parser.add_argument("--H", type=int, default=4)
    parser.add_argument("--HV", type=int, default=4, help="KDA value/gate heads. Must be a multiple of H.")
    parser.add_argument("--K", type=int, default=64)
    parser.add_argument("--V", type=int, default=64)
    parser.add_argument("--BT", type=int, default=64, choices=(32, 64))
    parser.add_argument("--BC", type=int, default=16)
    parser.add_argument("--dtype", choices=("float16", "bfloat16"), default="float16")
    return parser.parse_args()


def cell_col_index(ref: str) -> int:
    value = 0
    for ch in "".join(c for c in ref if c.isalpha()):
        value = value * 26 + ord(ch.upper()) - ord("A") + 1
    return value - 1


def parse_number(value: str | None):
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        as_float = float(text)
    except ValueError:
        return text
    return int(as_float) if as_float.is_integer() else as_float


def load_xlsx_rows(path: Path, sheet_name: str) -> list[tuple[int, list[object]]]:
    if not path.exists():
        raise FileNotFoundError(f"case workbook not found: {path}")
    with zipfile.ZipFile(path) as zf:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in zf.namelist():
            root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
            for si in root.findall(".//{*}si"):
                shared.append("".join(t.text or "" for t in si.findall(".//{*}t")))

        workbook = ET.fromstring(zf.read("xl/workbook.xml"))
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        rid_to_target = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rels.findall(".//{*}Relationship")}
        sheet_rid = None
        for sheet in workbook.findall(".//{*}sheet"):
            if sheet.attrib.get("name") == sheet_name:
                sheet_rid = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                break
        if sheet_rid is None:
            raise ValueError(f"sheet {sheet_name!r} not found in {path}")

        target = rid_to_target[sheet_rid]
        sheet_path = target.lstrip("/") if target.startswith("/") else str(PurePosixPath("xl") / target)
        sheet_xml = ET.fromstring(zf.read(sheet_path))
        rows: list[tuple[int, list[object]]] = []
        for row in sheet_xml.findall(".//{*}sheetData/{*}row"):
            row_num = int(row.attrib.get("r", len(rows) + 1))
            values: dict[int, object] = {}
            max_col = -1
            for cell in row.findall("{*}c"):
                col = cell_col_index(cell.attrib["r"])
                max_col = max(max_col, col)
                cell_type = cell.attrib.get("t")
                if cell_type == "inlineStr":
                    value = "".join(t.text or "" for t in cell.findall(".//{*}t"))
                else:
                    raw = cell.find("{*}v")
                    text = raw.text if raw is not None else None
                    if cell_type == "s" and text is not None:
                        value = shared[int(text)]
                    elif cell_type == "b" and text is not None:
                        value = text == "1"
                    else:
                        value = parse_number(text)
                values[col] = value
            rows.append((row_num, [values.get(i) for i in range(max_col + 1)]))
        return rows


def normalize_dtype(value: object) -> str:
    text = str(value).strip().lower()
    mapping = {"fp16": "float16", "float16": "float16", "bf16": "bfloat16", "bfloat16": "bfloat16"}
    if text not in mapping:
        raise ValueError(f"unsupported dtype in xlsx: {value!r}")
    return mapping[text]


def parse_bool(value: object) -> bool | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"", "n.a.", "na", "none", "null", "-"}:
        return None
    if text in {"true", "1", "yes", "y", "是"}:
        return True
    if text in {"false", "0", "no", "n", "否"}:
        return False
    raise ValueError(f"unsupported bool value in xlsx: {value!r}")


def parse_cu_seqlens(config: object) -> tuple[int, ...] | None:
    if config is None:
        return None
    text = str(config).strip()
    if not text:
        return None

    # ① 显式空值：cu_seqlens=None / cu_seqlens = none / cu_seqlens: null / nan
    if re.search(r"cu_seqlens\s*[=:]\s*(none|null|nan)\b", text, re.IGNORECASE):
        return None

    # ② 不含 cu_seqlens 关键字就当没有配置（大小写不敏感）
    if "cu_seqlens" not in text.lower():
        return None

    # ③ 解析列表 / range 两种写法
    values: list[int] = []
    match = re.search(r"cu_seqlens\s*=\s*\[([^\]]*)\]", text)
    if match:
        values = [int(x.strip()) for x in match.group(1).split(",") if x.strip()]
    for start, stop in re.findall(r"range\((\d+)\s*,\s*(\d+)\)", text):
        values.extend(range(int(start), int(stop)))
    if not values:
        raise ValueError(f"cannot parse cu_seqlens from: {text}")
    return tuple(values)


def int_field(row: dict[str, object], name: str) -> int:
    value = row.get(name)
    if value is None or str(value).strip() == "":
        raise ValueError(f"missing {name} in row {row}")
    return int(value)


def load_cases_from_xlsx(path: Path) -> list[BenchCase]:
    rows = load_xlsx_rows(path, "启用用例汇总")
    header: list[str] | None = None
    cases: list[BenchCase] = []
    for row_num, values in rows:
        if values and values[0] == "全局序号":
            header = [str(v).strip() if v is not None else "" for v in values]
            continue
        if header is None or not values or values[0] is None:
            continue
        row = {header[i]: values[i] if i < len(values) else None for i in range(len(header))}
        op = str(row.get("算子名称（用户）") or "").strip()
        if op not in OPS:
            continue
        d = int_field(row, "D")
        cases.append(BenchCase(
            op=op,
            op_index=int_field(row, "算子内序号"),
            label=str(row.get("Case 标签") or "").strip(),
            layout=str(row.get("布局") or "Padded").strip(),
            B=int_field(row, "B"),
            T=int_field(row, "T"),
            H=int_field(row, "H"),
            HV=int_field(row, "HV"),
            K=d,
            V=d,
            BT=int_field(row, "chunk_size"),
            BC=16,
            dtype=normalize_dtype(row.get("dtype")),
            safe_gate=parse_bool(row.get("safe_gate")),
            disable_recompute=bool(parse_bool(row.get("disable_recompute")) or False),
            cu_seqlens=parse_cu_seqlens(row.get("cu_seqlens / 附加配置")),
            config=str(row.get("cu_seqlens / 附加配置") or "").strip(),
            source_row=row_num,
        ))
    if not cases:
        raise ValueError(f"no benchmark cases found in {path}")
    return cases


def selected_cases(args: argparse.Namespace) -> list[BenchCase]:
    if args.single_case:
        cases = [BenchCase(args.op, 1, "single", "Padded", args.B, args.T, args.H, args.HV, args.K, args.V, args.BT, args.BC, args.dtype, args.safe_gate)]
    else:
        cases = load_cases_from_xlsx(args.case_xlsx)
    if args.op_filter:
        wanted_ops = {x.strip() for x in args.op_filter.split(",") if x.strip()}
        cases = [case for case in cases if case.op in wanted_ops]
    if args.case_filter:
        wanted = {x.strip() for x in args.case_filter.split(",") if x.strip()}
        cases = [case for case in cases if {case.case_id, case.label, str(case.op_index), f"{case.op_index:02d}"} & wanted]
    if not cases:
        raise ValueError("No cases selected.")
    return cases


def percentile(values: list[float], q: float) -> float:
    xs = sorted(values)
    idx = min(len(xs) - 1, max(0, int(round((len(xs) - 1) * q))))
    return xs[idx]


def time_cuda(fn: Callable[[], object], warmup: int, iters: int) -> list[float]:
    for _ in range(warmup):
        fn()
    torch.cuda.synchronize()
    times: list[float] = []
    for _ in range(iters):
        start = torch.cuda.Event(enable_timing=True)
        end = torch.cuda.Event(enable_timing=True)
        start.record()
        fn()
        end.record()
        torch.cuda.synchronize()
        times.append(start.elapsed_time(end))
    return times


def summarize(op: str, case: BenchCase, times: list[float], note: str) -> BenchResult:
    return BenchResult(op, case.case_id, case.shape_desc(), sum(times) / len(times), percentile(times, 0.50), percentile(times, 0.90), min(times), max(times), note)


def validate_case(case: BenchCase) -> None:
    assert case.HV % case.H == 0, f"{case.case_id}: HV must be a multiple of H"
    assert case.BT % case.BC == 0, f"{case.case_id}: BT must be a multiple of BC"
    assert case.BT in (32, 64), f"{case.case_id}: BT must be 32 or 64"
    assert case.BC == 16, f"{case.case_id}: this KDA script expects BC=16"
    if case.cu_seqlens is not None:
        assert case.B == 1, f"{case.case_id}: packed varlen cases should use B=1"
        assert case.cu_seqlens[0] == 0, f"{case.case_id}: cu_seqlens must start from 0"
        assert case.cu_seqlens[-1] == case.T, f"{case.case_id}: cu_seqlens[-1] must equal T"
        assert all(a < b for a, b in zip(case.cu_seqlens, case.cu_seqlens[1:])), f"{case.case_id}: cu_seqlens must increase"


def make_chunk_indices(case: BenchCase, device: torch.device):
    if case.cu_seqlens is None:
        return None, None
    from fla.ops.utils import prepare_chunk_indices
    cu_seqlens = torch.tensor(case.cu_seqlens, device=device, dtype=torch.long)
    return cu_seqlens, prepare_chunk_indices(cu_seqlens, case.BT)


def make_inputs(case: BenchCase, args: argparse.Namespace):
    validate_case(case)
    torch.manual_seed(args.seed)
    torch.cuda.set_device(args.device)
    device = torch.device(f"cuda:{args.device}")
    dtype = getattr(torch, case.dtype)
    B, T, H, HV, K, V, BT, BC = case.B, case.T, case.H, case.HV, case.K, case.V, case.BT, case.BC
    cu_seqlens, chunk_indices = make_chunk_indices(case, device)
    NT = chunk_indices.shape[0] if chunk_indices is not None else triton.cdiv(T, BT)

    q = torch.randn(B, T, H, K, device=device, dtype=dtype) * 0.1
    k = torch.randn(B, T, H, K, device=device, dtype=dtype) * 0.1
    v = torch.randn(B, T, HV, V, device=device, dtype=dtype) * 0.1
    do = torch.randn(B, T, HV, V, device=device, dtype=dtype) * 0.1
    g_dtype = torch.float32 if case.op == TOKEN_PARALLEL_OP or "gk/Aqk/Akk=float32" in case.config else dtype
    g_raw = torch.randn(B, T, HV, K, device=device, dtype=torch.float32) * 0.01
    g = torch.cumsum(g_raw, dim=1).to(g_dtype)
    beta = torch.sigmoid(torch.randn(B, T, HV, device=device, dtype=torch.float32)).to(dtype)
    mat_dtype = torch.float32 if case.op == TOKEN_PARALLEL_OP or "Aqk/Akk=float32" in case.config else dtype
    Aqk = torch.randn(B, T, HV, BT, device=device, dtype=mat_dtype) * 0.01
    Akk = torch.zeros(B, T, HV, BT, device=device, dtype=mat_dtype)
    Akkd = torch.empty(B, T, HV, BC, device=device, dtype=torch.float32)
    h = torch.randn(B, NT, HV, K, V, device=device, dtype=dtype) * 0.01
    dq = torch.zeros(B, T, HV, K, device=device, dtype=torch.float32)
    dk = torch.zeros(B, T, HV, K, device=device, dtype=torch.float32)
    db = torch.zeros(B, T, HV, device=device, dtype=torch.float32)
    dg = torch.zeros(B, T, HV, K, device=device, dtype=torch.float32)
    dh = torch.randn(B, NT, HV, K, V, device=device, dtype=dtype) * 0.01
    dv = torch.randn(B, T, HV, V, device=device, dtype=dtype) * 0.01
    v_new = torch.randn_like(v)
    dAqk = torch.randn(B, T, HV, BT, device=device, dtype=torch.float32) * 0.01
    dAkk = torch.randn(B, T, HV, BT, device=device, dtype=torch.float32) * 0.01
    return dict(q=q, k=k, v=v, do=do, g=g, beta=beta, Aqk=Aqk, Akk=Akk, Akkd=Akkd, h=h,
                dq=dq, dk=dk, db=db, dg=dg, dh=dh, dv=dv, v_new=v_new, dAqk=dAqk, dAkk=dAkk,
                NT=NT, scale=K ** -0.5, cu_seqlens=cu_seqlens, chunk_indices=chunk_indices)


def build_benchmark(case: BenchCase, args: argparse.Namespace, x: dict[str, object]):
    from fla.ops.gla.chunk import chunk_gla_fwd_kernel_o
    from fla.ops.kda.chunk_bwd import chunk_kda_bwd_dAv, chunk_kda_bwd_intra, chunk_kda_bwd_wy_dqkg_fused
    from fla.ops.kda.chunk_intra import chunk_kda_fwd_kernel_inter_solve_fused, chunk_kda_fwd_kernel_intra_sub_chunk
    from fla.ops.kda.chunk_intra_token_parallel import chunk_kda_fwd_intra_token_parallel
    from fla.utils import IS_GATHER_SUPPORTED

    B, T, H, HV, K, V, BT, BC = case.B, case.T, case.H, case.HV, case.K, case.V, case.BT, case.BC
    NT, scale = x["NT"], x["scale"]

    # Per-kernel BK/BV rules (VERIFIED against FLA source):
    #   - chunk_gla_fwd_kernel_o             : BK/BV in autotune configs {} -> DO NOT pass.
    #   - chunk_kda_fwd_kernel_inter_solve_fused : BK is REQUIRED constexpr -> MUST pass BK_INTER.
    #   - chunk_kda_fwd_kernel_intra_sub_chunk   : BK is REQUIRED constexpr -> MUST pass BK_INTRA.
    #   - chunk_kda_bwd_kernel_dAv / wy_dqkg     : wrappers pass BK/BV themselves.
    BK_INTRA = triton.next_power_of_2(K)
    BK_INTER = 32 if K <= 32 else 64
    NC = triton.cdiv(BT, BC)
    q, k, v, do = x["q"], x["k"], x["v"], x["do"]
    g, beta = x["g"], x["beta"]
    Aqk, Akk, Akkd, h = x["Aqk"], x["Akk"], x["Akkd"], x["h"]
    dq, dk, db, dg = x["dq"], x["dk"], x["db"], x["dg"]
    dh, dv, v_new = x["dh"], x["dv"], x["v_new"]
    dAqk, dAkk = x["dAqk"], x["dAkk"]
    cu_seqlens, chunk_indices = x["cu_seqlens"], x["chunk_indices"]

    # safe_gate resolution:
    #   --no-safe-gate  -> force False (workaround for Triton 3.2 b_gn bug)
    #   --safe-gate     -> force True
    #   otherwise       -> use xlsx value
    if args.no_safe_gate:
        use_safe_gate = False
    else:
        use_safe_gate = args.safe_gate or bool(case.safe_gate)

    def run_gla_o():
        o = torch.empty_like(v)

        def grid(meta):
            return (triton.cdiv(V, meta["BV"]), NT, B * HV)

        chunk_gla_fwd_kernel_o[grid](q=q, v=v, g=g, h=h, o=o, A=Aqk,
            cu_seqlens=cu_seqlens, chunk_indices=chunk_indices, scale=scale, T=T,
            H=H, HV=HV, K=K, V=V, BT=BT,
            STATE_V_FIRST=False, USE_GRAPH=False)
        return o

    def run_kda_dav():
        return chunk_kda_bwd_dAv(q=q, k=k, v=v, do=do, A=Aqk, scale=scale,
            cu_seqlens=cu_seqlens, chunk_size=BT, chunk_indices=chunk_indices, use_graph=False)

    def run_kda_bwd_intra():
        return chunk_kda_bwd_intra(q=q, k=k, g=g, beta=beta, dAqk=dAqk, dAkk=dAkk,
            dq=dq, dk=dk, db=db, dg=dg, cu_seqlens=cu_seqlens,
            chunk_indices=chunk_indices, chunk_size=BT, safe_gate=use_safe_gate, use_graph=False)

    def run_kda_wy_dqkg():
        return chunk_kda_bwd_wy_dqkg_fused(q=q, k=k, v=v, v_new=v_new, g=g, beta=beta, A=Akk,
            h=h, do=do, dh=dh, dv=dv, scale=scale, state_v_first=False,
            cu_seqlens=cu_seqlens, chunk_size=BT, chunk_indices=chunk_indices, use_graph=False)

    def run_inter_solve_fused():
        grid = (NT, B * HV)
        # BK is a REQUIRED constexpr here -> must be passed.
        chunk_kda_fwd_kernel_inter_solve_fused[grid](
            q=q, k=k, g=g, beta=beta, Aqk=Aqk, Akkd=Akkd, Akk=Akk,
            scale=scale, cu_seqlens=cu_seqlens, chunk_indices=chunk_indices, T=T,
            H=H, HV=HV, K=K, BT=BT, BC=BC, NC=NC, BK=BK_INTER,
            USE_SAFE_GATE=use_safe_gate, USE_GRAPH=False)
        return Akk

    def run_intra_sub_chunk():
        grid = (NT, NC, B * HV)
        # BK is a REQUIRED constexpr here -> must be passed.
        chunk_kda_fwd_kernel_intra_sub_chunk[grid](
            q=q, k=k, g=g, beta=beta, Aqk=Aqk, Akk=Akkd,
            scale=scale, cu_seqlens=cu_seqlens, chunk_indices=chunk_indices, T=T,
            H=H, HV=HV, K=K, BT=BT, BC=BC, BK=BK_INTRA,
            USE_GATHER=IS_GATHER_SUPPORTED, USE_GRAPH=False)
        return Akkd

    def run_token_parallel():
        return chunk_kda_fwd_intra_token_parallel(q=q, k=k, gk=g, beta=beta, Aqk=Aqk, Akk=Akkd,
            scale=scale, cu_seqlens=cu_seqlens, chunk_size=BT, sub_chunk_size=BC, use_graph=False)

    benches = {
        "chunk_gla_fwd_kernel_o": (run_gla_o, "direct Triton kernel"),
        "chunk_kda_bwd_dAv": (run_kda_dav, "Python wrapper, includes output allocation"),
        "chunk_kda_bwd_intra": (run_kda_bwd_intra, "Python wrapper, includes temporary allocation and db sum"),
        "chunk_kda_bwd_wy_dqkg_fused": (run_kda_wy_dqkg, "Python wrapper, includes output allocation"),
        "chunk_kda_fwd_intra_inter_solve_fused": (run_inter_solve_fused, "direct Triton kernel: actual name chunk_kda_fwd_kernel_inter_solve_fused"),
        "chunk_kda_fwd_intra_intra_sub_chunk": (run_intra_sub_chunk, "direct Triton kernel: actual name chunk_kda_fwd_kernel_intra_sub_chunk"),
        "chunk_kda_fwd_intra_token_parrallel": (run_token_parallel, "Python wrapper: actual spelling chunk_kda_fwd_intra_token_parallel"),
    }
    return case.op, *benches[case.op]


def prefill_diagonal(case: BenchCase, x: dict[str, object]) -> None:
    if case.op != INTER_SOLVE_OP:
        return
    from fla.ops.kda.chunk_intra_token_parallel import chunk_kda_fwd_intra_token_parallel
    chunk_kda_fwd_intra_token_parallel(q=x["q"], k=x["k"], gk=x["g"], beta=x["beta"], Aqk=x["Aqk"], Akk=x["Akkd"],
        scale=x["scale"], cu_seqlens=x["cu_seqlens"], chunk_size=case.BT, sub_chunk_size=case.BC, use_graph=False)
    torch.cuda.synchronize()


def print_case_header(args: argparse.Namespace, cases: list[BenchCase]) -> None:
    print("Benchmark config: " + ", ".join([f"device=cuda:{args.device}", f"warmup={args.warmup}", f"iters={args.iters}",
        f"num_cases={len(cases)}", f"case_xlsx={args.case_xlsx}"]))
    if args.no_safe_gate:
        print("[warn] --no-safe-gate is active: safe_gate forced to False for all cases.")
    print("\nCases:")
    for op in OPS:
        op_cases = [case for case in cases if case.op == op]
        if not op_cases:
            continue
        print(f"  {op}: {len(op_cases)} cases")
        for case in op_cases:
            print(f"    row={case.source_row:02d} case={case.case_id}: {case.shape_desc()}")
    print()


def print_result_header() -> None:
    print(f"{'op':42s} {'case':>18s} {'mean(ms)':>10s} {'p50':>10s} {'p90':>10s} {'min':>10s} {'max':>10s}  shape / note")
    print("-" * 166)


def print_dry_run(cases: list[BenchCase]) -> None:
    print(f"Loaded {len(cases)} cases.")
    counts = {op: 0 for op in OPS}
    for case in cases:
        validate_case(case)
        counts[case.op] += 1
    for op in OPS:
        print(f"{op}: {counts[op]} cases")
    print()
    for case in cases:
        seq_tail = f", cu_seqlens={case.cu_seqlens[:4]}...{case.cu_seqlens[-3:]}" if case.cu_seqlens else ""
        print(f"row={case.source_row:02d} {case.op:42s} {case.case_id:18s} {case.shape_desc()}{seq_tail}")


# --------------------------------------------------------------------------- #
# Result export
# --------------------------------------------------------------------------- #

def _fmt_cu_seqlens(case: BenchCase | None) -> str:
    if case is None or case.cu_seqlens is None:
        return ""
    seqs = case.cu_seqlens
    if len(seqs) <= 8:
        return str(list(seqs))
    return f"{list(seqs[:4])}...{list(seqs[-2:])} (len={len(seqs)})"


def _save_results_to_csv(path: Path, cases: list[BenchCase],
                         all_results: list[BenchResult],
                         failed: list[tuple[BenchCase, str]]) -> None:
    import csv
    case_by_key = {(c.op, c.case_id): c for c in cases}
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["op", "case_id", "layout", "B", "T", "H", "HV", "K", "V",
                    "BT", "BC", "dtype", "safe_gate", "cu_seqlens",
                    "mean_ms", "p50_ms", "p90_ms", "min_ms", "max_ms",
                    "status", "error", "note"])
        for r in all_results:
            c = case_by_key.get((r.op, r.case_id))
            w.writerow([r.op, r.case_id,
                        c.layout if c else "", c.B if c else "", c.T if c else "",
                        c.H if c else "", c.HV if c else "", c.K if c else "",
                        c.V if c else "", c.BT if c else "", c.BC if c else "",
                        c.dtype if c else "",
                        "" if c is None or c.safe_gate is None else bool(c.safe_gate),
                        _fmt_cu_seqlens(c),
                        r.mean_ms, r.p50_ms, r.p90_ms, r.min_ms, r.max_ms,
                        "OK", "", r.note])
        for c, err in failed:
            w.writerow([c.op, c.case_id, c.layout, c.B, c.T, c.H, c.HV, c.K, c.V,
                        c.BT, c.BC, c.dtype,
                        "" if c.safe_gate is None else bool(c.safe_gate),
                        _fmt_cu_seqlens(c),
                        "", "", "", "", "",
                        "FAILED", err, ""])
    print(f"\n[info] openpyxl unavailable -> results saved to CSV: {path}")


def save_results_to_xlsx(path: Path, args: argparse.Namespace,
                         cases: list[BenchCase],
                         all_results: list[BenchResult],
                         failed: list[tuple[BenchCase, str]],
                         wall_seconds: float) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        _save_results_to_csv(path.with_suffix(".csv"), cases, all_results, failed)
        return

    import datetime as _dt

    case_by_key = {(c.op, c.case_id): c for c in cases}
    wb = Workbook()

    # ---------------- Sheet 1: per-case results ---------------- #
    ws = wb.active
    ws.title = "results"
    headers = [
        "op", "case_id", "label", "layout", "B", "T", "H", "HV", "K", "V",
        "BT", "BC", "dtype", "safe_gate", "cu_seqlens",
        "mean_ms", "p50_ms", "p90_ms", "min_ms", "max_ms",
        "status", "error", "note", "source_row",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    def _case_base(c: BenchCase | None):
        if c is None:
            return ["", "", "", "", None, None, None, None, None, None,
                    None, None, "", "", ""]
        return [
            c.op, c.case_id, c.label, c.layout,
            c.B, c.T, c.H, c.HV, c.K, c.V,
            c.BT, c.BC, c.dtype,
            "" if c.safe_gate is None else bool(c.safe_gate),
            _fmt_cu_seqlens(c),
        ]

    for r in all_results:
        c = case_by_key.get((r.op, r.case_id))
        ws.append(_case_base(c) + [
            r.mean_ms, r.p50_ms, r.p90_ms, r.min_ms, r.max_ms,
            "OK", "", r.note, c.source_row if c else None,
        ])
    for c, err in failed:
        ws.append(_case_base(c) + [
            None, None, None, None, None,
            "FAILED", err, "", c.source_row,
        ])

    # ---------------- Sheet 2: summary per op ---------------- #
    ws2 = wb.create_sheet("summary_by_op")
    ws2.append([
        "op", "cases_ok", "cases_failed",
        "avg_mean_ms", "avg_p50_ms",
        "best_p50_ms", "worst_p50_ms",
        "best_min_ms", "worst_min_ms",
    ])
    for cell in ws2[1]:
        cell.font = Font(bold=True)

    by_op: dict[str, list[BenchResult]] = {}
    for r in all_results:
        by_op.setdefault(r.op, []).append(r)
    failed_by_op: dict[str, int] = {}
    for c, _ in failed:
        failed_by_op[c.op] = failed_by_op.get(c.op, 0) + 1

    for op in OPS:
        results = by_op.get(op, [])
        n_failed = failed_by_op.get(op, 0)
        if not results and n_failed == 0:
            continue
        if results:
            avg_mean = sum(r.mean_ms for r in results) / len(results)
            avg_p50 = sum(r.p50_ms for r in results) / len(results)
            best_p50 = min(r.p50_ms for r in results)
            worst_p50 = max(r.p50_ms for r in results)
            best_min = min(r.min_ms for r in results)
            worst_min = max(r.min_ms for r in results)
        else:
            avg_mean = avg_p50 = best_p50 = worst_p50 = best_min = worst_min = None
        ws2.append([op, len(results), n_failed,
                    avg_mean, avg_p50, best_p50, worst_p50, best_min, worst_min])

    # ---------------- Sheet 3: run configuration ---------------- #
    ws3 = wb.create_sheet("config")
    ws3.append(["key", "value"])
    ws3["A1"].font = Font(bold=True)
    ws3["B1"].font = Font(bold=True)
    ws3.append(["timestamp", _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws3.append(["device", f"cuda:{args.device}"])
    ws3.append(["warmup", args.warmup])
    ws3.append(["iters", args.iters])
    ws3.append(["seed", args.seed])
    ws3.append(["case_xlsx", str(args.case_xlsx)])
    ws3.append(["op_filter", args.op_filter])
    ws3.append(["case_filter", args.case_filter])
    ws3.append(["single_case", bool(args.single_case)])
    ws3.append(["no_safe_gate", bool(args.no_safe_gate)])
    ws3.append(["total_cases", len(cases)])
    ws3.append(["success", len(all_results)])
    ws3.append(["failed", len(failed)])
    ws3.append(["wall_seconds", round(wall_seconds, 3)])

    # ---------------- Column widths ---------------- #
    for sheet in (ws, ws2, ws3):
        for col_cells in sheet.iter_cols():
            col_letter = get_column_letter(col_cells[0].column)
            max_len = 0
            for c in col_cells:
                if c.value is None:
                    continue
                max_len = max(max_len, len(str(c.value)))
            sheet.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 60)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"\n[info] results saved to Excel: {path}")


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #

def main() -> None:
    args = parse_args()
    cases = selected_cases(args)
    if args.dry_run:
        print_dry_run(cases)
        return
    load_gpu_modules()
    if not torch.cuda.is_available():
        raise RuntimeError("CUDA is not available. Please run this script on a GPU machine, or use --dry-run for logic validation.")

    import time as _time

    print_case_header(args, cases)
    print_result_header()
    t_start = _time.perf_counter()
    all_results: list[BenchResult] = []
    failed: list[tuple[BenchCase, str]] = []
    for case in cases:
        x = make_inputs(case, args)
        prefill_diagonal(case, x)
        op, fn, note = build_benchmark(case, args, x)
        try:
            times = time_cuda(fn, args.warmup, args.iters)
            result = summarize(op, case, times, note)
            all_results.append(result)
            print(f"{result.op:42s} {result.case_id:>18s} {result.mean_ms:10.4f} {result.p50_ms:10.4f} {result.p90_ms:10.4f} {result.min_ms:10.4f} {result.max_ms:10.4f}  {result.shape} / {result.note}")
        except Exception as exc:
            torch.cuda.synchronize()
            err_msg = f"{type(exc).__name__}: {exc}"
            failed.append((case, err_msg))
            print(f"{op:42s} {case.case_id:>18s} {'FAILED':>10s}  {case.shape_desc()} / {err_msg}")
        finally:
            del x
            torch.cuda.empty_cache()
    wall_seconds = _time.perf_counter() - t_start

    if all_results:
        print("\nSummary by op:")
        by_op: dict[str, list[BenchResult]] = {}
        for result in all_results:
            by_op.setdefault(result.op, []).append(result)
        print(f"{'op':42s} {'cases':>6s} {'avg_p50(ms)':>13s} {'best_p50(ms)':>13s} {'worst_p50(ms)':>14s} {'best_min(ms)':>13s}")
        print("-" * 110)
        for op in OPS:
            results = by_op.get(op, [])
            if results:
                avg_p50 = sum(r.p50_ms for r in results) / len(results)
                best_p50 = min(r.p50_ms for r in results)
                worst_p50 = max(r.p50_ms for r in results)
                best_min = min(r.min_ms for r in results)
                print(f"{op:42s} {len(results):6d} {avg_p50:13.4f} {best_p50:13.4f} {worst_p50:14.4f} {best_min:13.4f}")

    if not args.no_save:
        save_results_to_xlsx(args.output_xlsx, args, cases, all_results, failed, wall_seconds)


if __name__ == "__main__":
    os.environ.setdefault("TRITON_PRINT_AUTOTUNING", "0")
    main()