#!/usr/bin/env python3
"""Generate the operator-specific pytest_v2 batch-consistency workbooks."""

import argparse
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


FIXED_TIME = datetime(2026, 7, 27, tzinfo=timezone.utc)
FIXED_ZIP_TIME = (2026, 7, 27, 0, 0, 0)
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
PRESSURE_FILL = PatternFill("solid", fgColor="FFF2CC")

SMLA_HEADERS = [
    "testcase_name", "layout_q", "layout_kv", "q_type", "ori_kv_type",
    "cmp_kv_type", "B", "S1", "S2", "T1", "T2", "T3", "N1", "N2", "D",
    "K", "block_num1", "block_num2", "block_size1", "block_size2", "seqused_q",
    "cu_seqlens_q", "cu_seqlens_ori_kv", "cu_seqlens_cmp_kv",
    "seqused_ori_kv", "seqused_cmp_kv", "cmp_residual_kv", "softmax_scale",
    "cmp_ratio", "ori_mask_mode", "cmp_mask_mode", "ori_win_left", "ori_win_right",
    "ori_kv_topk_mode", "cmp_kv_topk_mode", "ori_sparse_indices_mode",
    "cmp_sparse_indices_mode", "actlen_mode", "template_mode", "q_datarange",
    "ori_kv_datarange", "cmp_kv_datarange", "qk_equal_len", "return_softmax_lse",
    "consistency_modes", "split_groups", "token_split_sizes", "mode_batch",
    "common_tokens", "derived_extra", "purpose",
]

MQSMLA_HEADERS = [
    "Testcase_Name", "layout_q", "layout_kv", "q_type", "ori_kv_type",
    "cmp_kv_type", "B", "S1", "S2", "N1", "N2", "D", "K", "block_num1",
    "block_num2", "block_size1", "block_size2", "softmax_scale", "cmp_ratio",
    "ori_mask_mode", "cmp_mask_mode", "ori_win_left", "ori_win_right", "quant_mode",
    "tile_size", "rope_head_dim", "template_run_mode", "actlen_mode", "S1EQS2",
    "topk_value_mode", "return_softmax_lse", "seqused_q", "cu_seqlens_q",
    "seqused_ori_kv", "seqused_cmp_kv", "cu_seqlens_ori_kv",
    "cu_seqlens_cmp_kv", "cmp_residual_kv", "q_datarange", "ori_kv_datarange",
    "cmp_kv_datarange", "consistency_modes", "split_groups", "token_split_sizes",
    "mode_batch", "common_tokens", "derived_extra", "purpose",
]

PRESSURE_HEADERS = [
    "candidate_id", "operator", "layout_q", "layout_kv", "B", "S1_or_T1", "S2",
    "N1", "N2", "target_modes", "status", "unlock_gate", "purpose", "source",
]


def _smla_base() -> Dict[str, Any]:
    return {
        "layout_q": "BSND", "layout_kv": "BSND", "q_type": "torch.bfloat16",
        "ori_kv_type": "torch.bfloat16", "cmp_kv_type": "torch.bfloat16",
        "B": 2, "S1": 4, "S2": 8, "T1": 8, "T2": None, "T3": None,
        "N1": 4, "N2": 1, "D": 8, "K": None, "block_num1": None,
        "block_num2": None, "block_size1": 4, "block_size2": 4,
        "seqused_q": "[4,4]", "cu_seqlens_q": "[0,4,8]",
        "cu_seqlens_ori_kv": None, "cu_seqlens_cmp_kv": None,
        "seqused_ori_kv": "[8,8]", "seqused_cmp_kv": None,
        "cmp_residual_kv": None, "softmax_scale": 0.35355339, "cmp_ratio": None,
        "ori_mask_mode": 0, "cmp_mask_mode": 0, "ori_win_left": 127,
        "ori_win_right": 0, "ori_kv_topk_mode": "full", "cmp_kv_topk_mode": "full",
        "ori_sparse_indices_mode": "full", "cmp_sparse_indices_mode": "full",
        "actlen_mode": "full", "template_mode": "SWA", "q_datarange": "[-1,1]",
        "ori_kv_datarange": "[-1,1]", "cmp_kv_datarange": "[-1,1]",
        "qk_equal_len": True, "return_softmax_lse": False,
        "consistency_modes": "reorder,split,token-split,shape-change",
        "split_groups": "0;1", "token_split_sizes": "2,2", "mode_batch": 0,
        "common_tokens": 3, "derived_extra": 2,
    }


def smla_rows() -> List[Dict[str, Any]]:
    base = _smla_base()
    rows = []
    rows.append({
        **base,
        "testcase_name": "smla_bsnd_nomask_all_modes",
        "purpose": "Runnable CPU baseline: BSND no-mask coverage for all four named modes.",
    })
    rows.append({
        **base,
        "testcase_name": "smla_bsnd_rightdown_modes34",
        "ori_mask_mode": 3,
        "consistency_modes": "token-split,shape-change",
        "split_groups": None,
        "purpose": "Runnable CPU relation check: position-sensitive KV alignment for modes 3/4.",
    })
    rows.append({
        **base,
        "testcase_name": "smla_tnd_nomask_all_modes",
        "layout_q": "TND", "layout_kv": "TND", "T2": 16,
        "cu_seqlens_ori_kv": "[0,8,16]",
        "purpose": "Runnable CPU baseline: TND prefix-sum transforms for all four named modes.",
    })
    rows.append({
        **base,
        "testcase_name": "smla_pa_band_all_modes",
        "layout_kv": "PA_BBND", "ori_mask_mode": 4, "derived_extra": 0,
        "purpose": "Runnable CPU baseline: PA block-table transforms with a band mask; mode 4 shrinks within mapped capacity.",
    })
    return rows


def _mqsmla_base() -> Dict[str, Any]:
    return {
        "layout_q": "BSND", "layout_kv": "PA_BBND", "q_type": "torch.bfloat16",
        "ori_kv_type": "torch.float8_e4m3fn", "cmp_kv_type": "torch.float8_e4m3fn",
        "B": 2, "S1": 4, "S2": 8, "N1": 4, "N2": 1, "D": 128, "K": 0,
        "block_num1": None, "block_num2": None, "block_size1": 4, "block_size2": 4,
        "softmax_scale": 0.08838835, "cmp_ratio": 1, "ori_mask_mode": 4,
        "cmp_mask_mode": 0, "ori_win_left": 127, "ori_win_right": 0,
        "quant_mode": 1, "tile_size": 64, "rope_head_dim": 64,
        "template_run_mode": "SWA", "actlen_mode": "full", "S1EQS2": False,
        "topk_value_mode": 1, "return_softmax_lse": False, "seqused_q": "[4,4]",
        "cu_seqlens_q": "[0,4,8]", "seqused_ori_kv": "[8,8]",
        "seqused_cmp_kv": None, "cu_seqlens_ori_kv": "[0,8,16]",
        "cu_seqlens_cmp_kv": None, "cmp_residual_kv": None,
        "q_datarange": "[-1,1]", "ori_kv_datarange": "[-1,1]",
        "cmp_kv_datarange": "[-1,1]",
        "consistency_modes": "reorder,split,token-split,shape-change",
        "split_groups": "0;1", "token_split_sizes": "2,2", "mode_batch": 0,
        "common_tokens": 3, "derived_extra": 0,
    }


def mqsmla_rows() -> List[Dict[str, Any]]:
    base = _mqsmla_base()
    rows = []
    rows.append({
        **base,
        "Testcase_Name": "mqsmla_pa_band_all_modes",
        "purpose": "Runnable CPU baseline: PA quantized KV, band mask, and all four named modes; mode 4 stays within mapped capacity.",
    })
    rows.append({
        **base,
        "Testcase_Name": "mqsmla_bsnd_nomask_all_modes",
        "layout_kv": "BSND", "ori_mask_mode": 0, "derived_extra": 2,
        "purpose": "Runnable CPU baseline: BSND quantized KV and all four named modes.",
    })
    rows.append({
        **base,
        "Testcase_Name": "mqsmla_bsnd_rightdown_modes34",
        "layout_kv": "BSND", "ori_mask_mode": 3, "derived_extra": 2,
        "consistency_modes": "token-split,shape-change", "split_groups": None,
        "purpose": "Runnable CPU relation check: position-sensitive quantized KV alignment for modes 3/4.",
    })
    rows.append({
        **base,
        "Testcase_Name": "mqsmla_tnd_nomask_all_modes",
        "layout_q": "TND", "layout_kv": "TND", "ori_mask_mode": 0,
        "derived_extra": 2,
        "purpose": "Runnable CPU baseline: TND prefix-sum transforms for quantized KV and all four named modes.",
    })
    return rows


def pressure_rows(operator: str) -> List[Dict[str, Any]]:
    prefix = operator.lower()
    return [
        {
            "candidate_id": f"{prefix}_s2_4096_tile128", "operator": operator,
            "layout_q": "BSND", "layout_kv": "BSND", "B": 8, "S1_or_T1": 4,
            "S2": 4096, "N1": 64, "N2": 1, "target_modes": "reorder,split",
            "status": "NOT_EXECUTED_CANDIDATE",
            "unlock_gate": "runtime capability SUPPORTED, then minimal metadata/main-op smoke PASS",
            "purpose": "Exact 4096/32=128 reduction-boundary candidate.",
            "source": "STC scheduling intent; full executable fields require owner confirmation",
        },
        {
            "candidate_id": f"{prefix}_s2_5120_tile256", "operator": operator,
            "layout_q": "BSND", "layout_kv": "BSND", "B": 8, "S1_or_T1": 4,
            "S2": 5120, "N1": 64, "N2": 1, "target_modes": "reorder,split",
            "status": "NOT_EXECUTED_CANDIDATE",
            "unlock_gate": "runtime capability SUPPORTED, then resource estimate and short matrix PASS",
            "purpose": "Greater-than-128 and divisible boundary: 5120/32=160, aligned tile 256.",
            "source": "Corrected STC 014 category; not a non-divisible case",
        },
        {
            "candidate_id": f"{prefix}_n1_3", "operator": operator,
            "layout_q": "TND", "layout_kv": "TND", "B": 8, "S1_or_T1": 32,
            "S2": 1024, "N1": 3, "N2": 1, "target_modes": "reorder,token-split",
            "status": "NOT_EXECUTED_CANDIDATE",
            "unlock_gate": "owner confirms N means N1 and N1/N2 grouping is supported",
            "purpose": "Low-head-count scheduling candidate.",
            "source": "STC 025; original row does not define whether N is N1 or N2",
        },
        {
            "candidate_id": f"{prefix}_n1_125_pa", "operator": operator,
            "layout_q": "TND", "layout_kv": "PA_BBND", "B": 8, "S1_or_T1": 32,
            "S2": 1024, "N1": 125, "N2": 1, "target_modes": "reorder,split",
            "status": "NOT_EXECUTED_CANDIDATE",
            "unlock_gate": "owner confirms head semantics and PA capacity; runtime capability SUPPORTED",
            "purpose": "High-head-count plus paged-KV scheduling candidate.",
            "source": "STC 024/027; original rows have incomplete arrays and block mapping",
        },
    ]


def _append_dict_rows(sheet, headers: Iterable[str], rows: Iterable[Dict[str, Any]]) -> None:
    header_list = list(headers)
    sheet.append(header_list)
    for row in rows:
        unknown = sorted(set(row) - set(header_list))
        if unknown:
            raise ValueError(f"Unknown columns for {sheet.title}: {unknown}")
        sheet.append([row.get(name) for name in header_list])


def _style_table(sheet, pressure: bool = False) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PRESSURE_FILL if pressure else HEADER_FILL
        cell.font = Font(bold=True) if pressure else HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column in range(1, sheet.max_column + 1):
        values = [str(sheet.cell(row, column).value or "") for row in range(1, sheet.max_row + 1)]
        width = min(max(max(map(len, values)) + 2, 10), 54)
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _guide_rows(operator: str, case_sheet: str) -> List[List[str]]:
    return [
        ["item", "value"],
        ["operator", operator],
        ["runnable sheet", f"{case_sheet}: consumed by consistency_prepare and CPU-only."],
        ["mode syntax", "Use canonical names: reorder, split, token-split, shape-change. Numeric aliases are rejected."],
        ["pressure sheet", "PressureCandidates is documentation only; no row in it has been executed or passed."],
        ["NPU status", "A CPU prepare PASS proves parsing, golden generation, transforms, semantic oracle, and PT round-trip only."],
        ["unlock order", "capability SUPPORTED -> minimal metadata/main-op smoke -> four-mode matrix -> pressure candidates"],
    ]


def _save_reproducible(workbook: Workbook, path: Path) -> None:
    raw_path = path.with_name(f".{path.name}.raw")
    normalized_path = path.with_name(f".{path.name}.normalized")
    try:
        workbook.save(raw_path)
        with ZipFile(raw_path, "r") as source, ZipFile(
            normalized_path, "w", compression=ZIP_DEFLATED, compresslevel=9
        ) as target:
            for name in sorted(source.namelist()):
                original = source.getinfo(name)
                info = ZipInfo(name, FIXED_ZIP_TIME)
                info.compress_type = ZIP_DEFLATED
                info.external_attr = original.external_attr
                info.create_system = original.create_system
                target.writestr(info, source.read(name))
        normalized_path.replace(path)
    finally:
        raw_path.unlink(missing_ok=True)
        normalized_path.unlink(missing_ok=True)


def write_workbook(path: Path, operator: str, case_sheet: str, headers, rows) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "ops-transformer pytest_v2"
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME

    cases = workbook.create_sheet(case_sheet)
    _append_dict_rows(cases, headers, rows)
    _style_table(cases)

    pressure = workbook.create_sheet("PressureCandidates")
    _append_dict_rows(pressure, PRESSURE_HEADERS, pressure_rows(operator))
    _style_table(pressure, pressure=True)

    guide = workbook.create_sheet("Guide")
    for row in _guide_rows(operator, case_sheet):
        guide.append(row)
    _style_table(guide)

    path.parent.mkdir(parents=True, exist_ok=True)
    _save_reproducible(workbook, path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--repo-root",
        type=Path,
        default=Path(__file__).resolve().parents[2],
        help="ops-transformer_0716 checkout root",
    )
    args = parser.parse_args()
    root = args.repo_root.resolve()
    outputs = [
        (
            root / "attention/sparse_flash_mla/tests/pytest_v2/excel/smla_batch_consistency_cases.xlsx",
            "SMLA", "SMLA", SMLA_HEADERS, smla_rows(),
        ),
        (
            root / "attention/mixed_quant_sparse_flash_mla/tests/pytest_v2/excel/mqsmla_batch_consistency_cases.xlsx",
            "MQSMLA", "MQSMLA", MQSMLA_HEADERS, mqsmla_rows(),
        ),
    ]
    for path, operator, case_sheet, headers, rows in outputs:
        write_workbook(path, operator, case_sheet, headers, rows)
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
