#!/usr/bin/env python3
"""Convert aclnnGroupedMatmulV5 XLSX cases to the TTK GMM CSV schema.

The ACLNN workbook stores twelve inputs in API order.  TTK's grouped_matmul
case format uses nine inputs in this order:

    x, weight, bias, scale, offset, antiquant_scale, antiquant_offset,
    group_list, per_token_scale

Only the first tensor shape/dtype/format from each ACLNN tensor-list input is
needed because the source workbook's ``single_*`` sheets contain one tensor
per list.
"""

from __future__ import annotations

import argparse
import ast
import csv
import os
import re
import sys
import tempfile
from pathlib import Path
from typing import Any, Iterable, Sequence

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover - depends on the runtime image
    raise SystemExit("openpyxl is required: python3 -m pip install openpyxl") from exc


TTK_FIELDS = (
    "testcase_name",
    "op_name",
    "stc_inputs",
    "stc_ori_inputs",
    "stc_outputs",
    "stc_ori_outputs",
    "dyn_inputs",
    "dyn_ori_inputs",
    "dyn_outputs",
    "dyn_ori_outputs",
    "stc_input_formats",
    "dyn_input_formats",
    "dyn_input_ori_formats",
    "stc_input_ori_formats",
    "output_formats",
    "output_ori_formats",
    "stc_input_dtypes",
    "dyn_input_dtypes",
    "output_dtypes",
    "other_compilation_params",
    "dyn_input_ranges",
    "dyn_output_ranges",
    "dyn_input_as_list_distribution",
    "stc_input_as_list_distribution",
    "precision_tolerances",
    "input_data_ranges",
    "strict_precision_mode",
    "absolute_precision",
    "shape_check",
    "is_enabled",
)

# ACLNN input order -> TTK grouped_matmul input order.
INPUT_ORDER = (0, 1, 2, 4, 5, 6, 7, 3, 8)
OPTIONAL_TTK_INPUTS = frozenset((2, 3, 4, 5, 6, 8))
INPUT_LIST_DISTRIBUTION = (1, 1, 1, 1, 1, 1, 1, 0, 0, 1)
TRANSPOSE_FLAGS = {
    0: (False, False),
    1: (True, False),
    2: (False, True),
}
ATTRIBUTE_NAMES = {
    "grouplist": "group_list",
    "transpose": "transpose",
    "grouptype": "grouptype",
    "splititem": "split_item",
    "hasbias": "has_bias",
    "quantflag": "quant_flag",
    "a8w8flag": "a8w8_flag",
    "grouplisttype": "grouplisttype",
    "acttype": "acttype",
    "pta": "pta",
}

INPUT_DTYPE_ALIASES = {
    "fp16": "float16",
    "float16": "float16",
    "bf16": "bfloat16",
    "bfloat16": "bfloat16",
    "fp32": "float",
    "float32": "float",
}
OUTPUT_DTYPE_ALIASES = {
    **INPUT_DTYPE_ALIASES,
    "fp32": "float32",
    "float32": "float32",
}


def parse_literal(value: Any, field: str, sheet: str, row_number: int) -> Any:
    if not isinstance(value, str):
        return value
    try:
        return ast.literal_eval(value)
    except (SyntaxError, ValueError) as exc:
        raise ValueError(
            f"{sheet}!{row_number}: {field} is not a valid Python literal: {exc}"
        ) from exc


def require_int(value: Any, field: str, sheet: str, row_number: int) -> int:
    if isinstance(value, bool) or not isinstance(value, int):
        raise ValueError(
            f"{sheet}!{row_number}: {field} must be an integer, got {value!r}"
        )
    return value


def first_list_item(value: Any, field: str, sheet: str, row_number: int) -> Any:
    if not isinstance(value, list) or len(value) != 1:
        raise ValueError(
            f"{sheet}!{row_number}: {field} must contain exactly one item, got {value!r}"
        )
    return value[0]


def shape_tuple(value: Any, field: str, sheet: str, row_number: int) -> tuple[int, ...]:
    if not isinstance(value, (list, tuple)):
        raise ValueError(
            f"{sheet}!{row_number}: {field} must be a shape list, got {value!r}"
        )
    shape = tuple(value)
    if any(isinstance(dim, bool) or not isinstance(dim, int) for dim in shape):
        raise ValueError(
            f"{sheet}!{row_number}: {field} contains a non-integer dimension: {value!r}"
        )
    return shape


def is_placeholder_shape(shape: Sequence[int]) -> bool:
    return tuple(shape) == (1,)


def map_dtype(value: Any, output: bool = False) -> str:
    dtype = str(value).strip()
    aliases = OUTPUT_DTYPE_ALIASES if output else INPUT_DTYPE_ALIASES
    return aliases.get(dtype.lower(), dtype)


def infer_output_dtype(case_name: str, sheet: str, row_number: int) -> str:
    dtype_pattern = r"(bf16|bfloat16|fp16|float16|fp32|float32|int8|int32)"
    for pattern in (
        rf"output_{dtype_pattern}(?:_|$)",
        rf"MXFP8_{dtype_pattern}(?:_|$)",
    ):
        match = re.search(pattern, case_name, flags=re.IGNORECASE)
        if match:
            return map_dtype(match.group(1), output=True)
    raise ValueError(
        f"{sheet}!{row_number}: output dtype is empty and cannot be inferred "
        f"from case_name {case_name!r}"
    )


def tuple_text(value: Iterable[Any]) -> str:
    return repr(tuple(value))


def collect_attributes(record: dict[str, Any], sheet: str, row_number: int) -> dict[str, Any]:
    attrs: dict[str, Any] = {}
    for index in range(11):
        suffix = "" if index == 0 else f".{index}"
        name = record.get(f"attr_name{suffix}")
        if not name:
            continue
        raw_name = str(name)
        canonical_name = ATTRIBUTE_NAMES.get(
            raw_name.replace("_", "").lower(), raw_name
        )
        if canonical_name in attrs:
            raise ValueError(
                f"{sheet}!{row_number}: duplicate attribute {canonical_name!r}"
            )
        attrs[canonical_name] = parse_literal(
            record.get(f"attr_value{suffix}"), canonical_name, sheet, row_number
        )
    return attrs


def is_required_optional_input(
    ttk_index: int, quant_flag: int, a8w8_flag: int
) -> bool:
    # MX (quant_flag=3) passes scalar placeholders for offset and antiquant
    # tensors.  Other modes represent unused optional tensors as shape (0,).
    if quant_flag == 3 and ttk_index in (3, 4, 5, 6, 8):
        return True
    if ttk_index == 3:
        return quant_flag == 4 or a8w8_flag in (1, 2, 4)
    if ttk_index == 8:
        return quant_flag == 4 or a8w8_flag in (1, 4)
    return False


def derive_output_shape(
    x_shape: tuple[int, ...],
    weight_shape: tuple[int, ...],
    group_count: int,
    group_type: int,
    transpose_x: bool,
    transpose_weight: bool,
    sheet: str,
    row_number: int,
) -> tuple[int, ...]:
    if len(x_shape) != 2:
        raise ValueError(
            f"{sheet}!{row_number}: cannot derive output from x shape {x_shape!r}"
        )
    m_value = x_shape[1] if transpose_x else x_shape[0]

    # Weight-NZ cases already use a five-dimensional physical shape, so K/N
    # cannot be recovered reliably.  Their workbook row must provide output.
    if len(weight_shape) not in (2, 3):
        raise ValueError(
            f"{sheet}!{row_number}: output is empty and weight shape "
            f"{weight_shape!r} is not ND"
        )
    n_value = weight_shape[-2] if transpose_weight else weight_shape[-1]
    if group_type == 0:
        return (m_value, n_value)
    if group_type == 2:
        return (group_count, m_value, n_value)
    raise ValueError(
        f"{sheet}!{row_number}: cannot derive output for group_type={group_type}"
    )


def dynamic_shape(shape: Sequence[int]) -> tuple[int, ...]:
    return tuple(-1 for _ in shape)


def convert_record(
    record: dict[str, Any], sheet: str, row_number: int
) -> dict[str, str]:
    case_name = str(record.get("case_name") or "").strip()
    if not case_name:
        raise ValueError(f"{sheet}!{row_number}: case_name is empty")

    shapes = parse_literal(
        record.get("input_tensor_shape"), "input_tensor_shape", sheet, row_number
    )
    dtypes = parse_literal(
        record.get("input_tensor_dtype"), "input_tensor_dtype", sheet, row_number
    )
    formats = parse_literal(
        record.get("input_tensor_format"), "input_tensor_format", sheet, row_number
    )
    if not all(isinstance(item, list) and len(item) >= 9 for item in (shapes, dtypes, formats)):
        raise ValueError(
            f"{sheet}!{row_number}: input shape/dtype/format arrays must contain at least 9 items"
        )

    attrs = collect_attributes(record, sheet, row_number)
    required_attrs = (
        "group_list",
        "transpose",
        "grouptype",
        "split_item",
        "has_bias",
        "quant_flag",
        "a8w8_flag",
        "grouplisttype",
    )
    missing_attrs = [name for name in required_attrs if name not in attrs]
    if missing_attrs:
        raise ValueError(
            f"{sheet}!{row_number}: missing attributes: {', '.join(missing_attrs)}"
        )

    transpose = require_int(attrs["transpose"], "transpose", sheet, row_number)
    if transpose not in TRANSPOSE_FLAGS:
        raise ValueError(
            f"{sheet}!{row_number}: transpose must be 0, 1, or 2, got {transpose}"
        )
    transpose_x, transpose_weight = TRANSPOSE_FLAGS[transpose]
    group_type = require_int(attrs["grouptype"], "grouptype", sheet, row_number)
    split_item = require_int(attrs["split_item"], "split_item", sheet, row_number)
    has_bias = require_int(attrs["has_bias"], "has_bias", sheet, row_number)
    quant_flag = require_int(attrs["quant_flag"], "quant_flag", sheet, row_number)
    a8w8_flag = require_int(attrs["a8w8_flag"], "a8w8_flag", sheet, row_number)
    group_list_type = require_int(
        attrs["grouplisttype"], "grouplisttype", sheet, row_number
    )
    group_list = attrs["group_list"]
    if not isinstance(group_list, list) or any(
        isinstance(value, bool) or not isinstance(value, int) for value in group_list
    ):
        raise ValueError(
            f"{sheet}!{row_number}: group_list must be a list of integers"
        )

    stc_inputs: list[tuple[int, ...]] = []
    stc_input_dtypes: list[str] = []
    stc_input_formats: list[str] = []
    for ttk_index, aclnn_index in enumerate(INPUT_ORDER):
        if aclnn_index == 3:
            # group_list is the only plain tensor; all other API inputs are
            # one-element tensor lists.
            raw_shape = shapes[aclnn_index]
            dtype = dtypes[aclnn_index]
            tensor_format = formats[aclnn_index]
        else:
            raw_shape = first_list_item(
                shapes[aclnn_index],
                f"input_tensor_shape[{aclnn_index}]",
                sheet,
                row_number,
            )
            dtype = first_list_item(
                dtypes[aclnn_index],
                f"input_tensor_dtype[{aclnn_index}]",
                sheet,
                row_number,
            )
            tensor_format = first_list_item(
                formats[aclnn_index],
                f"input_tensor_format[{aclnn_index}]",
                sheet,
                row_number,
            )
        shape = shape_tuple(
            raw_shape, f"input_tensor_shape[{aclnn_index}][0]", sheet, row_number
        )

        present = True
        if ttk_index == 2:
            present = has_bias != 0
        elif ttk_index in OPTIONAL_TTK_INPUTS and is_placeholder_shape(shape):
            present = is_required_optional_input(ttk_index, quant_flag, a8w8_flag)
        if not present:
            shape = (0,)
        stc_inputs.append(shape)
        stc_input_dtypes.append("float" if ttk_index == 2 and not present else map_dtype(dtype))
        stc_input_formats.append(str(tensor_format))

    # TTK always consumes group_list as int64.
    stc_input_dtypes[7] = "int64"

    try:
        output_shapes = parse_literal(
            record.get("output_tensor_shape"), "output_tensor_shape", sheet, row_number
        )
    except ValueError:
        # A few source rows have a missing closing bracket.  Their output is
        # recoverable from x/weight shapes and transpose/group attributes.
        output_shapes = []
    output_dtypes = parse_literal(
        record.get("output_tensor_dtype"), "output_tensor_dtype", sheet, row_number
    )
    output_formats = parse_literal(
        record.get("output_tensor_format"), "output_tensor_format", sheet, row_number
    )
    output_shape: tuple[int, ...] = ()
    if isinstance(output_shapes, list) and output_shapes and output_shapes[0]:
        first_output_shape = output_shapes[0]
        raw_output_shape = (
            first_output_shape
            if all(isinstance(dim, int) and not isinstance(dim, bool) for dim in first_output_shape)
            else first_list_item(
                first_output_shape, "output_tensor_shape[0]", sheet, row_number
            )
        )
        output_shape = shape_tuple(
            raw_output_shape, "output_tensor_shape[0][0]", sheet, row_number
        )
    if not output_shape or any(dim <= 0 for dim in output_shape):
        output_shape = derive_output_shape(
            stc_inputs[0],
            stc_inputs[1],
            len(group_list),
            group_type,
            transpose_x,
            transpose_weight,
            sheet,
            row_number,
        )
    if isinstance(output_dtypes, list) and output_dtypes and output_dtypes[0]:
        first_output_dtype = output_dtypes[0]
        output_dtype = map_dtype(
            first_list_item(
                first_output_dtype, "output_tensor_dtype[0]", sheet, row_number
            )
            if isinstance(first_output_dtype, list)
            else first_output_dtype,
            output=True,
        )
    else:
        output_dtype = infer_output_dtype(case_name, sheet, row_number)
    if isinstance(output_formats, list) and output_formats and output_formats[0]:
        first_output_format = output_formats[0]
        output_format = str(
            first_list_item(
                first_output_format, "output_tensor_format[0]", sheet, row_number
            )
            if isinstance(first_output_format, list)
            else first_output_format
        )
    else:
        output_format = "ND"

    dyn_inputs = (
        dynamic_shape(stc_inputs[0]),
        dynamic_shape(stc_inputs[1]),
        (-2,),
        (-2,),
        (-1,),
        (-1,),
        (-1,),
        (-1,),
        (-2,),
    )
    dyn_outputs = (dynamic_shape(output_shape),)
    compilation_params = {
        "split_item": split_item,
        "group_list_type": group_list_type,
        "group_type": group_type,
        "transpose_x": transpose_x,
        "transpose_weight": transpose_weight,
        "group_list_expect": group_list,
    }

    precision = parse_literal(
        record.get("precision_tolerance"), "precision_tolerance", sheet, row_number
    )
    try:
        precision_pair = tuple(precision[0][:2])
    except (IndexError, TypeError) as exc:
        raise ValueError(
            f"{sheet}!{row_number}: precision_tolerance does not contain a pair"
        ) from exc

    stc_inputs_text = tuple_text(stc_inputs)
    stc_outputs_text = tuple_text((output_shape,))
    dyn_inputs_text = tuple_text(dyn_inputs)
    dyn_outputs_text = tuple_text(dyn_outputs)
    input_formats_text = tuple_text(stc_input_formats)
    input_dtypes_text = tuple_text(stc_input_dtypes)
    output_formats_text = tuple_text((output_format,))
    output_dtypes_text = tuple_text((output_dtype,))
    distribution_text = tuple_text(INPUT_LIST_DISTRIBUTION)

    return {
        "testcase_name": case_name,
        "op_name": "grouped_matmul",
        "stc_inputs": stc_inputs_text,
        "stc_ori_inputs": stc_inputs_text,
        "stc_outputs": stc_outputs_text,
        "stc_ori_outputs": stc_outputs_text,
        "dyn_inputs": dyn_inputs_text,
        "dyn_ori_inputs": dyn_inputs_text,
        "dyn_outputs": dyn_outputs_text,
        "dyn_ori_outputs": dyn_outputs_text,
        "stc_input_formats": input_formats_text,
        "dyn_input_formats": tuple_text(("ND",)),
        "dyn_input_ori_formats": input_formats_text,
        "stc_input_ori_formats": input_formats_text,
        "output_formats": output_formats_text,
        "output_ori_formats": output_formats_text,
        "stc_input_dtypes": input_dtypes_text,
        "dyn_input_dtypes": input_dtypes_text,
        "output_dtypes": output_dtypes_text,
        "other_compilation_params": repr(compilation_params),
        "dyn_input_ranges": "",
        "dyn_output_ranges": "",
        "dyn_input_as_list_distribution": distribution_text,
        "stc_input_as_list_distribution": distribution_text,
        "precision_tolerances": tuple_text((precision_pair,)),
        "input_data_ranges": "",
        "strict_precision_mode": "TRUE",
        "absolute_precision": "0.001",
        "shape_check": "TRUE",
        "is_enabled": "1",
    }


def resolve_sheets(workbook: Any, requested: Sequence[str]) -> list[str]:
    names = list(workbook.sheetnames)
    if requested == ["all"]:
        return names
    missing = [name for name in requested if name not in names]
    if missing:
        raise ValueError(
            f"sheet(s) not found: {', '.join(missing)}; available: {', '.join(names)}"
        )
    return list(requested)


def validate_reference_schema(path: Path) -> None:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        header = next(csv.reader(file), None)
    if tuple(header or ()) != TTK_FIELDS:
        raise ValueError(
            f"reference CSV header does not match the expected {len(TTK_FIELDS)}-column TTK schema"
        )


def convert_workbook(args: argparse.Namespace) -> tuple[int, list[str]]:
    if args.reference_csv:
        validate_reference_schema(args.reference_csv)

    workbook = openpyxl.load_workbook(args.input, read_only=True, data_only=True)
    try:
        sheet_names = resolve_sheets(workbook, args.sheet)
        converted: list[dict[str, str]] = []
        seen_names: dict[str, str] = {}
        for sheet_name in sheet_names:
            worksheet = workbook[sheet_name]
            rows = worksheet.iter_rows(values_only=True)
            try:
                headers = next(rows)
            except StopIteration as exc:
                raise ValueError(f"sheet {sheet_name!r} is empty") from exc
            if not headers or "case_name" not in headers:
                raise ValueError(f"sheet {sheet_name!r} has no case_name column")
            for row_number, values in enumerate(rows, start=2):
                record = dict(zip(headers, values))
                if not record.get("case_name"):
                    continue
                result = convert_record(record, sheet_name, row_number)
                previous_sheet = seen_names.get(result["testcase_name"])
                if previous_sheet and not args.allow_duplicate_names:
                    raise ValueError(
                        f"duplicate testcase_name {result['testcase_name']!r} in "
                        f"{previous_sheet} and {sheet_name}; select one sheet or use "
                        "--allow-duplicate-names"
                    )
                seen_names[result["testcase_name"]] = sheet_name
                converted.append(result)
    finally:
        workbook.close()

    if args.dry_run:
        return len(converted), sheet_names
    if args.output.exists() and not args.force:
        raise ValueError(f"output already exists: {args.output}; use --force to replace it")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{args.output.name}.", dir=args.output.parent
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        with temporary_path.open("w", encoding="utf-8-sig", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=TTK_FIELDS, lineterminator="\n")
            writer.writeheader()
            writer.writerows(converted)
        os.replace(temporary_path, args.output)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()
    return len(converted), sheet_names


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Convert ACLNN QuantGroupMatMul XLSX cases to TTK grouped_matmul CSV."
    )
    parser.add_argument("input", type=Path, help="source ACLNN .xlsx workbook")
    parser.add_argument("-o", "--output", type=Path, help="destination TTK .csv file")
    parser.add_argument(
        "--sheet",
        action="append",
        default=None,
        help=(
            "sheet to convert; repeat for multiple sheets, or use 'all' "
            "(default: single_aclnn)"
        ),
    )
    parser.add_argument(
        "--reference-csv",
        type=Path,
        help="optional existing TTK CSV whose 30-column header is validated",
    )
    parser.add_argument(
        "--allow-duplicate-names",
        action="store_true",
        help="allow duplicate testcase_name values when converting multiple sheets",
    )
    parser.add_argument("--force", action="store_true", help="replace an existing output")
    parser.add_argument(
        "--dry-run", action="store_true", help="validate and count cases without writing"
    )
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    args.input = args.input.resolve()
    args.sheet = args.sheet or ["single_aclnn"]
    if "all" in args.sheet and args.sheet != ["all"]:
        parser.error("--sheet all cannot be combined with another --sheet value")
    if args.reference_csv:
        args.reference_csv = args.reference_csv.resolve()
    if args.output is None:
        args.output = args.input.with_name(f"{args.input.stem}_ttk.csv")
    args.output = args.output.resolve()
    try:
        count, sheets = convert_workbook(args)
    except (OSError, ValueError, csv.Error) as exc:
        parser.error(str(exc))
    action = "validated" if args.dry_run else "converted"
    target = "" if args.dry_run else f", output: {args.output}"
    print(f"{action} {count} case(s) from {', '.join(sheets)}{target}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
