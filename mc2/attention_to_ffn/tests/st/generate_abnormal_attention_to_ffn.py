import os
import json
import copy

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

BASE_CONFIG = {
    "X": 1,
    "L": 1,
    "BS": 8,
    "H": 1024,
    "K": 8,
    "worldSize": 16,
    "moeExpertNum": 8,
    "sharedExpertNum": 1,
    "localmoeExpert": 1,
    "quantMode": 0,
    "syncFlag": 0,
    "ffnStartRankId": 0,
}

BASE_CONFIG["moeRankNum"] = BASE_CONFIG["moeExpertNum"] // BASE_CONFIG["localmoeExpert"]
BASE_CONFIG["sharedRankNum"] = BASE_CONFIG["sharedExpertNum"]
BASE_CONFIG["ffnRankNum"] = BASE_CONFIG["moeRankNum"] + BASE_CONFIG["sharedRankNum"]
BASE_CONFIG["attnRankNum"] = BASE_CONFIG["worldSize"] - BASE_CONFIG["ffnRankNum"]
BASE_CONFIG["totalExpertNum"] = BASE_CONFIG["moeExpertNum"] + BASE_CONFIG["sharedExpertNum"]
BASE_CONFIG["expertNumPerToken"] = BASE_CONFIG["K"] + BASE_CONFIG["sharedExpertNum"]
BASE_CONFIG["M"] = BASE_CONFIG["ffnRankNum"] * 2 + 1
BASE_CONFIG["HS"] = BASE_CONFIG["H"] + 128

INPUT_NAMES_ORDER = [
    "x", "sessionId", "microBatchId", "layerId",
    "expertIds", "expertRankTable", "scalesOptional", "activeMaskOptional",
    "worldSize", "ffnTokenInfoTableShape", "ffnTokenDataShape", "attnTokenInfoTableShape",
    "moeExpertNum", "quantMode", "syncFlag", "ffnStartRankId",
]

def _input_index(name):
    return INPUT_NAMES_ORDER.index(name)

def make_base_inputs(cfg):
    return [
        {"name": "x", "type": "tensor", "required": True, "dtype": "fp16",
         "shape": [cfg["X"], cfg["BS"], cfg["H"]], "range_values": [-1, 1]},
        {"name": "sessionId", "type": "tensor", "required": True, "dtype": "int32",
         "shape": [1], "range_values": [0]},
        {"name": "microBatchId", "type": "tensor", "required": True, "dtype": "int32",
         "shape": [1], "range_values": [0]},
        {"name": "layerId", "type": "tensor", "required": True, "dtype": "int32",
         "shape": [1], "range_values": [0]},
        {"name": "expertIds", "type": "tensor", "required": True, "dtype": "int32",
         "shape": [cfg["X"], cfg["BS"], cfg["K"]], "range_values": [0, 7]},
        {"name": "expertRankTable", "type": "tensor", "required": True, "dtype": "int32",
         "shape": [cfg["L"], cfg["totalExpertNum"], cfg["M"]], "range_values": [0, 1]},
        {"name": "scalesOptional", "type": "tensor", "required": True, "dtype": "fp32",
         "shape": [cfg["L"], cfg["totalExpertNum"], cfg["H"]], "range_values": [0, 0]},
        {"name": "activeMaskOptional", "type": "tensor", "required": True, "dtype": "bool",
         "shape": [cfg["X"], cfg["BS"]], "range_values": [True]},
        {"name": "worldSize", "type": "attr", "required": True, "dtype": "int64_t",
         "shape": None, "range_values": cfg["worldSize"]},
        {"name": "ffnTokenInfoTableShape", "type": "attr_tuple", "required": True,
         "dtype": "int32_t", "shape": None,
         "range_values": [cfg["attnRankNum"], 1, 2 + cfg["BS"] * cfg["expertNumPerToken"]]},
        {"name": "ffnTokenDataShape", "type": "attr_tuple", "required": True,
         "dtype": "int32_t", "shape": None,
         "range_values": [cfg["attnRankNum"], 1, cfg["BS"], cfg["expertNumPerToken"], cfg["H"]]},
        {"name": "attnTokenInfoTableShape", "type": "attr_tuple", "required": True,
         "dtype": "int32_t", "shape": None,
         "range_values": [1, cfg["BS"], cfg["expertNumPerToken"]]},
        {"name": "moeExpertNum", "type": "attr", "required": True, "dtype": "int64_t",
         "shape": None, "range_values": cfg["moeExpertNum"]},
        {"name": "quantMode", "type": "attr", "required": True, "dtype": "int64_t",
         "shape": None, "range_values": cfg["quantMode"]},
        {"name": "syncFlag", "type": "attr", "required": True, "dtype": "int64_t",
         "shape": None, "range_values": cfg["syncFlag"]},
        {"name": "ffnStartRankId", "type": "attr", "required": True, "dtype": "int64_t",
         "shape": None, "range_values": cfg["ffnStartRankId"]},
    ]

def make_base_case(case_id, cfg=None):
    if cfg is None:
        cfg = BASE_CONFIG
    return {
        "id": case_id,
        "default_seed": None,
        "name": "AttentionToFFN",
        "aclnn_name": "AttentionToFFN",
        "expected_error_msg": None,
        "api": "pytorch",
        "api_type": "execute_Attention2FFN",
        "aclnn_api_type": "pyaclnn_AttentionToFFN",
        "dist_api_type": "execute_Attention2FFN",
        "backward": False,
        "standard": {"acc": "default", "perf": "not_key"},
        "outputs": None,
        "inputs": make_base_inputs(cfg),
        "acl_json": "",
    }

def modify_input(case, input_name, **kwargs):
    for inp in case["inputs"]:
        if inp["name"] == input_name:
            for k, v in kwargs.items():
                inp[k] = v
            break
    return case

ABNORMAL_CASES_DEF = []

ABNORMAL_CASES_DEF.append({
    "category": "数据类型非法",
    "cases": [
        {"param": "x", "desc": "x不支持fp32", "modifications": {"dtype": "fp32"},
         "expected_error": "ACLNN_ERR_PARAM_INVALID"},
        {"param": "x", "desc": "x不支持int8", "modifications": {"dtype": "int8"},
         "expected_error": "ACLNN_ERR_PARAM_INVALID"},
        {"param": "sessionId", "desc": "sessionId不支持fp16", "modifications": {"dtype": "fp16"},
         "expected_error": "ACLNN_ERR_PARAM_INVALID"},
        {"param": "microBatchId", "desc": "microBatchId不支持fp16", "modifications": {"dtype": "fp16"},
         "expected_error": "ACLNN_ERR_PARAM_INVALID"},
        {"param": "layerId", "desc": "layerId不支持fp16", "modifications": {"dtype": "fp16"},
         "expected_error": "ACLNN_ERR_PARAM_INVALID"},
        {"param": "expertIds", "desc": "expertIds不支持fp16", "modifications": {"dtype": "fp16"},
         "expected_error": "ACLNN_ERR_PARAM_INVALID"},
        {"param": "expertRankTable", "desc": "expertRankTable不支持fp16", "modifications": {"dtype": "fp16"},
         "expected_error": "ACLNN_ERR_PARAM_INVALID"},
        {"param": "scalesOptional", "desc": "scalesOptional不支持fp16(应为fp32)",
         "modifications": {"dtype": "fp16"}, "extra_cfg": {"quantMode": 2, "is_attn2ffn_scales": 1},
         "expected_error": "ACLNN_ERR_PARAM_INVALID"},
    ]
})

ABNORMAL_CASES_DEF.append({
    "category": "维度非法",
    "cases": [
        {"param": "x", "desc": "X=2(当前版本只支持X=1)",
         "modifications": {"shape": [2, 8, 1024]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "x", "desc": "Bs=0(Bs应>0)",
         "modifications": {"shape": [1, 0, 1024]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "x", "desc": "H=512(H应>=1024)",
         "modifications": {"shape": [1, 8, 512]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "x", "desc": "H=10000(H应<=8192)",
         "modifications": {"shape": [1, 8, 10000]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "x", "desc": "Bs=600(Bs应<=512)",
         "modifications": {"shape": [1, 600, 1024]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "x", "desc": "x为2D(应为3D)",
         "modifications": {"shape": [1, 8]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "sessionId", "desc": "sessionId为2D(应为1D)",
         "modifications": {"shape": [1, 1]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "microBatchId", "desc": "microBatchId为2D(应为1D)",
         "modifications": {"shape": [1, 1]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "layerId", "desc": "layerId为2D(应为1D)",
         "modifications": {"shape": [1, 1]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "expertIds", "desc": "expertIds为2D(应为3D)",
         "modifications": {"shape": [8, 8]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "expertRankTable", "desc": "expertRankTable为2D(应为3D)",
         "modifications": {"shape": [1, 9]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
    ]
})

ABNORMAL_CASES_DEF.append({
    "category": "空Tensor非法",
    "cases": [
        {"param": "sessionId", "desc": "sessionId为空tensor(shape=[0])",
         "modifications": {"shape": [0]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "microBatchId", "desc": "microBatchId为空tensor(shape=[0])",
         "modifications": {"shape": [0]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "layerId", "desc": "layerId为空tensor(shape=[0])",
         "modifications": {"shape": [0]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "expertIds", "desc": "expertIds为空tensor(Bs=0)",
         "modifications": {"shape": [1, 0, 8]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "expertRankTable", "desc": "expertRankTable为空tensor(L=0)",
         "modifications": {"shape": [0, 9, 19]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
    ]
})

ABNORMAL_CASES_DEF.append({
    "category": "属性值非法",
    "cases": [
        {"param": "worldSize", "desc": "worldSize=1(<2)",
         "modifications": {"range_values": 1},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "worldSize", "desc": "worldSize=769(>768)",
         "modifications": {"range_values": 769},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "moeExpertNum", "desc": "moeExpertNum=0(<=0)",
         "modifications": {"range_values": 0},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "moeExpertNum", "desc": "moeExpertNum=1025(>1024)",
         "modifications": {"range_values": 1025},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "quantMode", "desc": "quantMode=1(静态量化不支持)",
         "modifications": {"range_values": 1},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "quantMode", "desc": "quantMode=3(>2)",
         "modifications": {"range_values": 3},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "quantMode", "desc": "quantMode=-1(<0)",
         "modifications": {"range_values": -1},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "syncFlag", "desc": "syncFlag=2(只支持0和1)",
         "modifications": {"range_values": 2},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "syncFlag", "desc": "syncFlag=-1(只支持0和1)",
         "modifications": {"range_values": -1},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "ffnStartRankId", "desc": "ffnStartRankId=-1(<0)",
         "modifications": {"range_values": -1},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "ffnStartRankId", "desc": "ffnStartRankId=16(>=worldSize)",
         "modifications": {"range_values": 16},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
    ]
})

ABNORMAL_CASES_DEF.append({
    "category": "空指针非法",
    "cases": [
        {"param": "x", "desc": "x为空指针(必选参数)",
         "attrs": {"isNullptr": [_input_index("x")]},
         "expected_error": "ACLNN_ERR_PARAM_NULLPTR"},
        {"param": "sessionId", "desc": "sessionId为空指针(必选参数)",
         "attrs": {"isNullptr": [_input_index("sessionId")]},
         "expected_error": "ACLNN_ERR_PARAM_NULLPTR"},
        {"param": "microBatchId", "desc": "microBatchId为空指针(必选参数)",
         "attrs": {"isNullptr": [_input_index("microBatchId")]},
         "expected_error": "ACLNN_ERR_PARAM_NULLPTR"},
        {"param": "layerId", "desc": "layerId为空指针(必选参数)",
         "attrs": {"isNullptr": [_input_index("layerId")]},
         "expected_error": "ACLNN_ERR_PARAM_NULLPTR"},
        {"param": "expertIds", "desc": "expertIds为空指针(必选参数)",
         "attrs": {"isNullptr": [_input_index("expertIds")]},
         "expected_error": "ACLNN_ERR_PARAM_NULLPTR"},
        {"param": "expertRankTable", "desc": "expertRankTable为空指针(必选参数)",
         "attrs": {"isNullptr": [_input_index("expertRankTable")]},
         "expected_error": "ACLNN_ERR_PARAM_NULLPTR"},
    ]
})

ABNORMAL_CASES_DEF.append({
    "category": "格式非法",
    "cases": [
        {"param": "x", "desc": "x格式为NZ(应为ND)",
         "attrs": {"isFmt": [_input_index("x")]},
         "expected_error": "ACLNN_ERR_PARAM_INVALID"},
        {"param": "expertIds", "desc": "expertIds格式为NZ(应为ND)",
         "attrs": {"isFmt": [_input_index("expertIds")]},
         "expected_error": "ACLNN_ERR_PARAM_INVALID"},
        {"param": "expertRankTable", "desc": "expertRankTable格式为NZ(应为ND)",
         "attrs": {"isFmt": [_input_index("expertRankTable")]},
         "expected_error": "ACLNN_ERR_PARAM_INVALID"},
    ]
})

ABNORMAL_CASES_DEF.append({
    "category": "Shape组合非法",
    "cases": [
        {"param": "ffnTokenInfoTableShape", "desc": "ffnTokenInfoTableShape长度不为3(传4个元素)",
         "modifications": {"range_values": [7, 1, 74, 1]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "ffnTokenDataShape", "desc": "ffnTokenDataShape长度不为5(传6个元素)",
         "modifications": {"range_values": [7, 1, 8, 9, 1024, 1]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "attnTokenInfoTableShape", "desc": "attnTokenInfoTableShape长度不为3(传4个元素)",
         "modifications": {"range_values": [1, 8, 9, 1]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "ffnTokenInfoTableShape", "desc": "ffnTokenInfoTableShape[0]!=ffnTokenDataShape[0]",
         "modifications": {"range_values": [5, 1, 74]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "ffnTokenInfoTableShape", "desc": "ffnTokenInfoTableShape[1]!=ffnTokenDataShape[1]",
         "modifications": {"range_values": [7, 2, 74]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "ffnTokenDataShape", "desc": "ffnTokenDataShape[2]!=bs(x.shape[1])",
         "modifications": {"range_values": [7, 1, 16, 9, 1024]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "attnTokenInfoTableShape", "desc": "attnTokenInfoTableShape[0]!=ffnTokenInfoTableShape[1]",
         "modifications": {"range_values": [2, 8, 9]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "attnTokenInfoTableShape", "desc": "attnTokenInfoTableShape[1]!=bs(x.shape[1])",
         "modifications": {"range_values": [1, 16, 9]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "attnTokenInfoTableShape", "desc": "attnTokenInfoTableShape[2]!=ffnTokenDataShape[3]",
         "modifications": {"range_values": [1, 8, 5]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "ffnTokenDataShape", "desc": "attentionWorkerNum=0(ffnTokenDataShape[0]=0)",
         "modifications": {"range_values": [0, 1, 8, 9, 1024]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "ffnTokenDataShape", "desc": "attentionWorkerNum>=worldSize(ffnTokenDataShape[0]=16)",
         "modifications": {"range_values": [16, 1, 8, 9, 1024]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
    ]
})

ABNORMAL_CASES_DEF.append({
    "category": "Optional参数shape组合非法",
    "cases": [
        {"param": "scalesOptional", "desc": "scales.shape[0]!=expertRankTable.shape[0]",
         "modifications": {"shape": [2, 9, 1024]},
         "extra_cfg": {"quantMode": 2, "is_attn2ffn_scales": 1},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "scalesOptional", "desc": "scales.shape[1]!=expertRankTable.shape[1]",
         "modifications": {"shape": [1, 5, 1024]},
         "extra_cfg": {"quantMode": 2, "is_attn2ffn_scales": 1},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "scalesOptional", "desc": "scales.shape[2]!=x.shape[2](H)",
         "modifications": {"shape": [1, 9, 512]},
         "extra_cfg": {"quantMode": 2, "is_attn2ffn_scales": 1},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "activeMaskOptional", "desc": "activeMask.shape[0]!=x.shape[0](X)",
         "modifications": {"shape": [2, 8]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
        {"param": "activeMaskOptional", "desc": "activeMask.shape[1]!=x.shape[1](Bs)",
         "modifications": {"shape": [1, 16]},
         "expected_error": "ACLNN_ERR_INNER_TILING_ERROR"},
    ]
})


def _get_doc_ref(category, param, desc):
    refs = {
        "数据类型非法": "文档参数说明-数据类型列；tiling CheckTensorDataType",
        "维度非法": "文档Shape变量约束表；tiling CheckAndSetShapeAttrs",
        "空Tensor非法": "文档Shape变量约束表(Bs>0等)；tiling校验",
        "属性值非法": "文档参数说明-使用说明列；tiling CheckAndSetAttrs",
        "空指针非法": "aclnn_attention_to_ffn.cpp CheckNullStatus；文档返回值ACLNN_ERR_PARAM_NULLPTR",
        "格式非法": "文档参数说明-数据格式列(ND)；def.cpp Format",
        "Shape组合非法": "tiling CheckAndSetShapeAttrs各一致性校验",
        "Optional参数shape组合非法": "tiling CheckInputForScales/CheckInputForActiveMask",
    }
    return refs.get(category, "")


def generate_xlsx(output_path, cases_def):
    if not HAS_OPENPYXL:
        print("[WARN] openpyxl not installed, skipping xlsx generation")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "非法场景枚举表"

    headers = ["序号", "非法维度", "参数名", "非法描述", "非法构造方式",
               "预期错误码", "文档约束依据"]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border

    col_widths = [6, 16, 22, 40, 40, 28, 40]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    row_idx = 2
    case_num = 1
    for category_def in cases_def:
        category = category_def["category"]
        for case in category_def["cases"]:
            param = case["param"]
            desc = case["desc"]
            expected = case.get("expected_error", "")

            if "attrs" in case:
                attrs = case["attrs"]
                if "isNullptr" in attrs:
                    construct = f"isNullptr={attrs['isNullptr']}"
                elif "isFmt" in attrs:
                    construct = f"isFmt={attrs['isFmt']}"
                elif "isNotContiguous" in attrs:
                    construct = f"isNotContiguous={attrs['isNotContiguous']}"
                else:
                    construct = str(attrs)
            elif "modifications" in case:
                mods = case["modifications"]
                parts = []
                for k, v in mods.items():
                    parts.append(f"{k}={v}")
                construct = ", ".join(parts)
            else:
                construct = ""

            doc_ref = _get_doc_ref(category, param, desc)

            row_data = [case_num, category, param, desc, construct, expected, doc_ref]
            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                cell.border = thin_border
            row_idx += 1
            case_num += 1

    wb.save(output_path)
    print(f"[OK] xlsx saved to {output_path}")


def generate_json(output_path, cases_def):
    all_cases = []
    case_id = 0

    for category_def in cases_def:
        category = category_def["category"]
        for case in category_def["cases"]:
            cfg = copy.deepcopy(BASE_CONFIG)
            if "extra_cfg" in case:
                cfg.update(case["extra_cfg"])

            json_case = make_base_case(case_id, cfg)

            if "modifications" in case:
                modify_input(json_case, case["param"], **case["modifications"])

            if "attrs" in case:
                for attr_key, attr_val in case["attrs"].items():
                    json_case["inputs"].append({
                        "name": attr_key,
                        "type": "attrs",
                        "required": False,
                        "dtype": "list_int64_t",
                        "shape": None,
                        "range_values": attr_val,
                    })

            if "extra_cfg" in case:
                for ek, ev in case["extra_cfg"].items():
                    found = False
                    for inp in json_case["inputs"]:
                        if inp["name"] == ek:
                            if inp["type"] == "attr":
                                inp["range_values"] = ev
                            found = True
                            break
                    if not found:
                        json_case["inputs"].append({
                            "name": ek,
                            "type": "attr",
                            "required": False,
                            "dtype": "int64_t",
                            "shape": None,
                            "range_values": ev,
                        })

            json_case["expected_error_msg"] = case.get("expected_error", "")
            all_cases.append(json_case)
            case_id += 1

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(all_cases, f, indent=4, ensure_ascii=False)
    print(f"[OK] json saved to {output_path} ({len(all_cases)} cases)")


if __name__ == "__main__":
    script_dir = os.path.dirname(os.path.abspath(__file__))
    xlsx_path = os.path.join(script_dir, "abnormal_cases.xlsx")
    json_path = os.path.join(script_dir, "abnormal_cases.json")

    generate_xlsx(xlsx_path, ABNORMAL_CASES_DEF)
    generate_json(json_path, ABNORMAL_CASES_DEF)

    print("\n=== 非法场景枚举汇总 ===")
    total = 0
    for category_def in ABNORMAL_CASES_DEF:
        cat = category_def["category"]
        cnt = len(category_def["cases"])
        total += cnt
        print(f"  {cat}: {cnt} 条")
    print(f"  总计: {total} 条")
