#!/usr/bin/env python3
"""
Engram算子测试用例生成器（更新版）
根据新要求重新设计：
1. 移除num_qps, HYBRID_MODE参数
2. 简化world_size：直接用2,4,8,16卡
3. fetch输出增加fp16, fp32支持
4. 新增performance sheet（仅fetch性能）
5. 新增exception sheet（失败用例）
"""

import os
import random
from typing import List, Dict, Any
from openpyxl import Workbook

# 参数定义（简化版）
HIDDEN_TIER = [64, 128, 256, 512, 1024, 2048, 4096]
ENTRIES_TIER = [4096, 8192, 16384, 32768, 65536, 131072, 262144, 524288]
NUM_TOKENS_TIER = [32, 64, 128, 256, 512, 1024, 2048, 4096, 8192]
WORLD_SIZE_TIER = [2, 4, 8, 16]
OUTPUT_DTYPE_TIER = ['bf16', 'fp16', 'fp32']
PHASE_TIER = ['write', 'fetch', 'write_fetch']


def generate_case_name(params: Dict[str, Any], case_id: int, sheet_type: str = 'level0') -> str:
    """生成符合MC2规范的case_name"""
    hidden = params.get('hidden', 256)
    num_entries = params.get('num_entries', 65536)
    num_tokens = params.get('num_tokens', 1024)
    world_size = params.get('world_size', 8)
    output_dtype = params.get('output_dtype', 'bf16')
    phase = params.get('phase', 'write_fetch')
    
    if sheet_type == 'performance':
        name_parts = [
            f"perf",
            f"h{hidden}",
            f"tok{num_tokens}",
            f"ws{world_size}",
            output_dtype,
            f"ID{case_id:03d}"
        ]
    elif sheet_type == 'exception':
        name_parts = [
            f"fail",
            phase,
            f"ID{case_id:03d}"
        ]
    else:
        name_parts = [
            f"h{hidden}",
            f"tok{num_tokens}",
            f"e{num_entries}",
            f"ws{world_size}",
            output_dtype,
            phase,
            f"ID{case_id:03d}"
        ]
    
    return "_".join(name_parts)


def create_case(params: Dict[str, Any], case_id: int, sheet_type: str = 'level0') -> Dict[str, Any]:
    """创建单个测试用例"""
    phase = params.get('phase', 'write_fetch')
    world_size = params.get('world_size', 8)
    num_entries = params.get('num_entries', 65536)
    num_tokens = params.get('num_tokens', 1024)
    hidden = params.get('hidden', 256)
    output_dtype = params.get('output_dtype', 'bf16')
    
    # 根据phase定义输入输出tensor
    if phase == 'write':
        input_shapes = [[num_entries, hidden]]
        input_dtypes = ['bf16']  # write输入总是bf16
        input_ranges = [(-5, 5)]
        output_dtypes = []
    elif phase == 'fetch':
        indices_max = num_entries * world_size - 1
        input_shapes = [[num_tokens]]
        input_dtypes = ['int32']
        input_ranges = [(0, indices_max)]
        output_dtypes = [output_dtype]  # fetch输出支持多种dtype
    else:  # write_fetch
        indices_max = num_entries * world_size - 1
        input_shapes = [[num_entries, hidden], [num_tokens]]
        input_dtypes = ['bf16', 'int32']
        input_ranges = [(-5, 5), (0, indices_max)]
        output_dtypes = [output_dtype]
    
    case = {
        'case_name': generate_case_name(params, case_id, sheet_type),
        'input_tensor_shape': input_shapes,
        'input_tensor_dtype': input_dtypes,
        'input_tensor_range': input_ranges,
        'output_tensor_dtype': output_dtypes,
        'hidden': hidden,
        'num_entries': num_entries,
        'num_tokens': num_tokens,
        'world_size': world_size,
        'output_dtype': output_dtype,
        'phase': phase,
        'seed': case_id,
        'graph_type': 0
    }
    
    return case


def generate_level0_cases() -> List[Dict[str, Any]]:
    """生成Level0基础覆盖用例（单因子覆盖，≤200条）"""
    cases = []
    case_id = 0
    
    # 基础参数
    base_params = {
        'hidden': 256,
        'num_entries': 65536,
        'num_tokens': 1024,
        'world_size': 8,
        'output_dtype': 'bf16',
        'phase': 'write_fetch'
    }
    
    # 1. HIDDEN_TIER覆盖
    for hidden in HIDDEN_TIER:
        params = base_params.copy()
        params['hidden'] = hidden
        cases.append(create_case(params, case_id))
        case_id += 1
    
    # 2. ENTRIES_TIER覆盖（选择代表性）
    for num_entries in [4096, 16384, 65536, 262144, 524288]:
        params = base_params.copy()
        params['num_entries'] = num_entries
        cases.append(create_case(params, case_id))
        case_id += 1
    
    # 3. NUM_TOKENS_TIER覆盖（选择代表性）
    for num_tokens in [32, 128, 512, 2048, 4096]:
        params = base_params.copy()
        params['num_tokens'] = num_tokens
        cases.append(create_case(params, case_id))
        case_id += 1
    
    # 4. WORLD_SIZE_TIER覆盖
    for world_size in WORLD_SIZE_TIER:
        params = base_params.copy()
        params['world_size'] = world_size
        cases.append(create_case(params, case_id))
        case_id += 1
    
    # 5. OUTPUT_DTYPE_TIER覆盖（fetch输出类型）
    for output_dtype in OUTPUT_DTYPE_TIER:
        params = base_params.copy()
        params['output_dtype'] = output_dtype
        params['phase'] = 'fetch'  # 仅fetch测试dtype
        cases.append(create_case(params, case_id))
        case_id += 1
    
    # 6. PHASE_TIER覆盖
    for phase in PHASE_TIER:
        params = base_params.copy()
        params['phase'] = phase
        if phase == 'write':
            params['output_dtype'] = ''
        cases.append(create_case(params, case_id))
        case_id += 1
    
    return cases[:200]


def generate_level1_cases() -> List[Dict[str, Any]]:
    """生成Level1两两组合用例（500~700条）"""
    cases = []
    case_id = 200
    
    # 1. HIDDEN × NUM_TOKENS组合
    for hidden in [128, 256, 512, 1024, 2048]:
        for num_tokens in [256, 512, 1024, 2048, 4096]:
            params = {
                'hidden': hidden,
                'num_entries': 65536,
                'num_tokens': num_tokens,
                'world_size': 8,
                'output_dtype': 'bf16',
                'phase': 'fetch'
            }
            cases.append(create_case(params, case_id))
            case_id += 1
    
    # 2. HIDDEN × WORLD_SIZE组合
    for hidden in [128, 256, 512, 1024]:
        for world_size in [2, 4, 8, 16]:
            params = {
                'hidden': hidden,
                'num_entries': 65536,
                'num_tokens': 1024,
                'world_size': world_size,
                'output_dtype': 'bf16',
                'phase': 'fetch'
            }
            cases.append(create_case(params, case_id))
            case_id += 1
    
    # 3. OUTPUT_DTYPE × HIDDEN组合
    for output_dtype in OUTPUT_DTYPE_TIER:
        for hidden in [128, 256, 512, 1024, 2048]:
            params = {
                'hidden': hidden,
                'num_entries': 65536,
                'num_tokens': 1024,
                'world_size': 8,
                'output_dtype': output_dtype,
                'phase': 'fetch'
            }
            cases.append(create_case(params, case_id))
            case_id += 1
    
    # 4. PHASE × HIDDEN组合
    for phase in PHASE_TIER:
        for hidden in [128, 256, 512, 1024]:
            params = {
                'hidden': hidden,
                'num_entries': 65536,
                'num_tokens': 1024,
                'world_size': 8,
                'output_dtype': 'bf16',
                'phase': phase
            }
            if phase == 'write':
                params['output_dtype'] = ''
            cases.append(create_case(params, case_id))
            case_id += 1
    
    # 5. WORLD_SIZE × OUTPUT_DTYPE组合
    for world_size in [2, 4, 8, 16]:
        for output_dtype in OUTPUT_DTYPE_TIER:
            params = {
                'hidden': 256,
                'num_entries': 65536,
                'num_tokens': 1024,
                'world_size': world_size,
                'output_dtype': output_dtype,
                'phase': 'fetch'
            }
            cases.append(create_case(params, case_id))
            case_id += 1
    
    # 6. 随机组合补充至500条
    while len(cases) < 500:
        params = {
            'hidden': random.choice(HIDDEN_TIER),
            'num_entries': random.choice(ENTRIES_TIER),
            'num_tokens': random.choice(NUM_TOKENS_TIER),
            'world_size': random.choice(WORLD_SIZE_TIER),
            'output_dtype': random.choice(OUTPUT_DTYPE_TIER),
            'phase': random.choice(PHASE_TIER)
        }
        if params['phase'] == 'write':
            params['output_dtype'] = ''
        cases.append(create_case(params, case_id))
        case_id += 1
    
    return cases[:700]


def generate_performance_cases() -> List[Dict[str, Any]]:
    """生成Performance性能用例（仅关注fetch接口）"""
    cases = []
    case_id = 0
    
    # 性能测试典型场景
    perf_configs = [
        # 小hidden测试
        {'hidden': 128, 'num_tokens': 4096, 'world_size': 8, 'output_dtype': 'bf16'},
        {'hidden': 128, 'num_tokens': 4096, 'world_size': 8, 'output_dtype': 'fp16'},
        {'hidden': 128, 'num_tokens': 4096, 'world_size': 8, 'output_dtype': 'fp32'},
        
        # 中hidden测试
        {'hidden': 256, 'num_tokens': 4096, 'world_size': 8, 'output_dtype': 'bf16'},
        {'hidden': 256, 'num_tokens': 4096, 'world_size': 8, 'output_dtype': 'fp16'},
        {'hidden': 256, 'num_tokens': 4096, 'world_size': 8, 'output_dtype': 'fp32'},
        
        # 大hidden测试
        {'hidden': 1024, 'num_tokens': 4096, 'world_size': 8, 'output_dtype': 'bf16'},
        {'hidden': 1024, 'num_tokens': 4096, 'world_size': 8, 'output_dtype': 'fp16'},
        {'hidden': 1024, 'num_tokens': 4096, 'world_size': 8, 'output_dtype': 'fp32'},
        
        # 不同world_size测试
        {'hidden': 256, 'num_tokens': 4096, 'world_size': 2, 'output_dtype': 'bf16'},
        {'hidden': 256, 'num_tokens': 4096, 'world_size': 4, 'output_dtype': 'bf16'},
        {'hidden': 256, 'num_tokens': 4096, 'world_size': 16, 'output_dtype': 'bf16'},
        
        # 不同num_tokens测试
        {'hidden': 256, 'num_tokens': 128, 'world_size': 8, 'output_dtype': 'bf16'},
        {'hidden': 256, 'num_tokens': 1024, 'world_size': 8, 'output_dtype': 'bf16'},
        {'hidden': 256, 'num_tokens': 8192, 'world_size': 8, 'output_dtype': 'bf16'},
        
        # 边界组合测试
        {'hidden': 64, 'num_tokens': 32, 'world_size': 2, 'output_dtype': 'bf16'},
        {'hidden': 4096, 'num_tokens': 8192, 'world_size': 16, 'output_dtype': 'fp32'},
    ]
    
    for config in perf_configs:
        params = {
            'hidden': config['hidden'],
            'num_entries': 65536,  # 固定entries
            'num_tokens': config['num_tokens'],
            'world_size': config['world_size'],
            'output_dtype': config['output_dtype'],
            'phase': 'fetch'  # 性能测试仅关注fetch
        }
        cases.append(create_case(params, case_id, 'performance'))
        case_id += 1
    
    # 增加更多组合（梯度测试）
    for hidden in [128, 256, 512, 1024, 2048, 4096]:
        params = {
            'hidden': hidden,
            'num_entries': 65536,
            'num_tokens': 4096,
            'world_size': 8,
            'output_dtype': 'bf16',
            'phase': 'fetch'
        }
        cases.append(create_case(params, case_id, 'performance'))
        case_id += 1
    
    for num_tokens in [32, 128, 512, 1024, 2048, 4096, 8192]:
        params = {
            'hidden': 256,
            'num_entries': 65536,
            'num_tokens': num_tokens,
            'world_size': 8,
            'output_dtype': 'bf16',
            'phase': 'fetch'
        }
        cases.append(create_case(params, case_id, 'performance'))
        case_id += 1
    
    return cases


def generate_exception_cases() -> List[Dict[str, Any]]:
    """生成Exception失败用例"""
    cases = []
    case_id = 0
    
    # 1. indices越界失败
    exception_configs = [
        {
            'desc': 'indices负数越界',
            'hidden': 256,
            'num_entries': 65536,
            'num_tokens': 1024,
            'world_size': 8,
            'output_dtype': 'bf16',
            'phase': 'fetch',
            'exception_type': 'indices_negative',
            'exception_detail': 'indices包含负数索引，预期失败'
        },
        {
            'desc': 'indices超大越界',
            'hidden': 256,
            'num_entries': 65536,
            'num_tokens': 1024,
            'world_size': 8,
            'output_dtype': 'bf16',
            'phase': 'fetch',
            'exception_type': 'indices_overflow',
            'exception_detail': 'indices >= world_size × num_entries，预期失败'
        },
        
        # 2. tensor参数错误
        {
            'desc': 'storage dtype错误',
            'hidden': 256,
            'num_entries': 65536,
            'num_tokens': 1024,
            'world_size': 8,
            'output_dtype': 'bf16',
            'phase': 'write',
            'exception_type': 'dtype_error',
            'exception_detail': 'storage dtype != bf16，预期失败'
        },
        {
            'desc': 'storage shape错误',
            'hidden': 256,
            'num_entries': 65536,
            'num_tokens': 1024,
            'world_size': 8,
            'output_dtype': 'bf16',
            'phase': 'write',
            'exception_type': 'shape_error',
            'exception_detail': 'storage shape != (num_entries, hidden)，预期失败'
        },
        
        # 3. 状态错误
        {
            'desc': 'buffer未构造',
            'hidden': 256,
            'num_entries': 65536,
            'num_tokens': 1024,
            'world_size': 8,
            'output_dtype': 'bf16',
            'phase': 'fetch',
            'exception_type': 'buffer_not_created',
            'exception_detail': 'buffer未构造直接调用API，预期失败'
        },
        {
            'desc': 'buffer已销毁',
            'hidden': 256,
            'num_entries': 65536,
            'num_tokens': 1024,
            'world_size': 8,
            'output_dtype': 'bf16',
            'phase': 'fetch',
            'exception_type': 'buffer_destroyed',
            'exception_detail': 'buffer.destroy()后调用API，预期失败'
        },
        
        # 4. 时序错误
        {
            'desc': 'fetch先于write',
            'hidden': 256,
            'num_entries': 65536,
            'num_tokens': 1024,
            'world_size': 8,
            'output_dtype': 'bf16',
            'phase': 'fetch',
            'exception_type': 'sequence_error',
            'exception_detail': 'engram_fetch先于engram_write，预期失败'
        },
        
        # 5. 并发错误
        {
            'desc': '并发fetch同一buffer',
            'hidden': 256,
            'num_entries': 65536,
            'num_tokens': 1024,
            'world_size': 8,
            'output_dtype': 'bf16',
            'phase': 'fetch',
            'exception_type': 'concurrent_error',
            'exception_detail': '同buffer并发fetch，last_request覆盖，预期失败'
        },
        
        # 6. device错误
        {
            'desc': 'tensor device不匹配',
            'hidden': 256,
            'num_entries': 65536,
            'num_tokens': 1024,
            'world_size': 8,
            'output_dtype': 'bf16',
            'phase': 'write_fetch',
            'exception_type': 'device_error',
            'exception_detail': 'tensor device不正确，预期失败'
        },
        
        # 7. 分布式错误
        {
            'desc': 'ProcessGroup未初始化',
            'hidden': 256,
            'num_entries': 65536,
            'num_tokens': 1024,
            'world_size': 8,
            'output_dtype': 'bf16',
            'phase': 'write_fetch',
            'exception_type': 'distributed_error',
            'exception_detail': 'ProcessGroup未初始化，预期失败'
        },
    ]
    
    for config in exception_configs:
        params = {
            'hidden': config['hidden'],
            'num_entries': config['num_entries'],
            'num_tokens': config['num_tokens'],
            'world_size': config['world_size'],
            'output_dtype': config['output_dtype'],
            'phase': config['phase']
        }
        case = create_case(params, case_id, 'exception')
        case['exception_type'] = config['exception_type']
        case['exception_detail'] = config['exception_detail']
        cases.append(case)
        case_id += 1
    
    return cases


def write_cases_to_sheet(ws, cases: List[Dict[str, Any]], columns: List[str]):
    """将用例写入sheet"""
    # 写入标题行
    for col_idx, col_name in enumerate(columns, 1):
        ws.cell(row=1, column=col_idx, value=col_name)
    
    # 写入数据行
    for row_idx, case in enumerate(cases, 2):
        for col_idx, col_name in enumerate(columns, 1):
            value = case.get(col_name, '')
            
            # 编码值
            if isinstance(value, list):
                encoded = str(value)
            elif isinstance(value, tuple):
                encoded = str([value]) if col_name == 'input_tensor_range' else str(value)
            elif isinstance(value, bool):
                encoded = '1' if value else '0'
            else:
                encoded = str(value)
            
            ws.cell(row=row_idx, column=col_idx, value=encoded)


def generate_xlsx(output_path: str):
    """生成xlsx文件"""
    wb = Workbook()
    
    # 删除默认sheet
    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']
    
    # 生成各sheet用例
    level0_cases = generate_level0_cases()
    level1_cases = generate_level1_cases()
    performance_cases = generate_performance_cases()
    exception_cases = generate_exception_cases()
    
    # 定义列顺序
    level_columns = [
        'case_name', 'input_tensor_shape', 'input_tensor_dtype', 'input_tensor_range',
        'output_tensor_dtype', 'hidden', 'num_entries', 'num_tokens',
        'world_size', 'output_dtype', 'phase', 'seed', 'graph_type'
    ]
    
    perf_columns = level_columns
    
    exception_columns = level_columns + ['exception_type', 'exception_detail']
    
    # 创建各sheet
    ws0 = wb.create_sheet('level0')
    write_cases_to_sheet(ws0, level0_cases, level_columns)
    
    ws1 = wb.create_sheet('level1')
    write_cases_to_sheet(ws1, level1_cases, level_columns)
    
    ws_perf = wb.create_sheet('performance')
    write_cases_to_sheet(ws_perf, performance_cases, perf_columns)
    
    ws_exc = wb.create_sheet('exception')
    write_cases_to_sheet(ws_exc, exception_cases, exception_columns)
    
    # 保存文件
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    
    print(f"Generated: {output_path}")
    print(f"  Level0: {len(level0_cases)} cases")
    print(f"  Level1: {len(level1_cases)} cases")
    print(f"  Performance: {len(performance_cases)} cases")
    print(f"  Exception: {len(exception_cases)} cases")


if __name__ == '__main__':
    output_path = 'mc2/result/excel/engram.xlsx'
    generate_xlsx(output_path)